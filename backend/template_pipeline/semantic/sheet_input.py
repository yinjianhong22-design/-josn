# -*- coding: utf-8 -*-

# ---------------------------------------------------------------------------
# 数据简表: 一次性匹配 JSON (用完即弃, 不落盘)
#
# 数据侧只回答两个问题: 这张表是什么 (title 路由模板), 有哪些字段和值
# (name 匹配模板索引; specified/measured[]/pass_rate 是要填的值)。
# 坐标/边框/字体一概不需要 —— 值该进哪几个格子由模板自己决定
# (index rows[].columns[列名].cells: 实测值列区子格按 x 序的 uuid 清单,
# measured 按序拆格; 规定值/合格率列保持单格)。
# 已填表照片/扫描件走同一出口: 识别+按行标签取值后产出同样的
# {title, fields}, 之后共用 match_data_fields。
# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# 数据简表: 一次性匹配 JSON (用完即弃, 不落盘)
# ---------------------------------------------------------------------------


def photo_to_sheet(data):
    """已填表照片 -> 数据简表 JSON {title, context, fields}。

    入口是 extract_page + assign_roles/extract_fields/build_index 的产物
    (与模板制作同一套识别), 出口与 parse_data_sheet 同构, 走同一
    match_data_fields 匹配。拍照件透视矫正是独立前置步骤, 这里假定
    正面清晰照片。"""
    import re

    def norm(s):
        return re.sub(r'\s+', '', str(s or ''))

    title_kw = ('记录表', '检验表', '检测报告', '评定表', '验收表', '检查表')
    ctx_kw = ('单位', '合同段', '人', '时间', '部位', '编号', '桩号', '日期')
    texts = sorted(data.get('texts') or [],
                   key=lambda t: (t.get('y') or 0, t.get('x') or 0))
    title = None
    for t in texts:
        s = norm(t.get('text'))
        if len(s) >= 6 and any(k in s for k in title_kw):
            if title is None or len(s) > len(title):
                title = s
    if title is None:
        # 标题被圈进格子时 (拍照件/截图整页成格, 标题不是游离文字),
        # 退而搜格内文字: 同一关键词过滤, 取最靠上的一格, 并按首个
        # 关键词截断去掉粘连的 "编号:..." 尾巴
        cand = []
        for c in data.get('cells') or []:
            s = norm(c.get('text'))
            if len(s) >= 6 and any(k in s for k in title_kw):
                cand.append((c.get('y') or 0, s))
        if cand:
            cand.sort(key=lambda p: p[0])
            s = cand[0][1]
            ends = [s.find(k) + len(k) for k in title_kw if k in s]
            title = s[:min(ends)] if ends else s

    idx = data.get('index') or {}
    cells = idx.get('cells') or {}

    def cell_text(u):
        return ((cells.get(u) or {}).get('text') or '').strip()

    def ctx_value(name, t):
        """整行共享格 "监理单位：… 编 号：JB12-x" -> 冒号后值段。"""
        m = re.match(r'^(.*?)[:：]\s*(\S.*)$', t or '')
        if m and norm(name) in norm(m.group(1)):
            return m.group(2).strip()
        return t

    context = {}
    for name, ent in (idx.get('header_fields') or {}).items():
        if not any(k in name for k in ctx_kw):
            continue
        u = ent.get('cell') if isinstance(ent, dict) else ent
        t = cell_text(u)
        if t:
            context[name] = ctx_value(name, t)
    for name, ents in (idx.get('by_label') or {}).items():
        if name in context or not any(k in name for k in ctx_kw):
            continue
        u = next((e.get('value') for e in ents if e.get('value')), None)
        t = cell_text(u) if u else ''
        if t:
            context[name] = ctx_value(name, t)

    fields = []
    for r in idx.get('rows') or []:
        f = {'name': (r.get('key') or '').strip()}
        if not f['name']:
            continue
        got = False
        for hn, col in (r.get('columns') or {}).items():
            nhn = norm(hn)
            if '规定值' in nhn or '允许偏差' in nhn or '标准值' in nhn:
                t = cell_text(col['cell'])
                if t:
                    f['specified'] = t
                    got = True
            elif '实测' in nhn:
                vals = []
                for u in (col.get('cells') or [col['cell']]):
                    vals.extend(_split_numbers(cell_text(u)))
                if vals:
                    f['measured'] = vals
                    got = True
            elif '合格率' in nhn or '通过率' in nhn:
                t = cell_text(col['cell']).replace('%', '').strip()
                try:
                    f['pass_rate'] = float(t) if '.' in t else int(t)
                    got = True
                except ValueError:
                    pass
        if got and (f.get('measured') or f.get('specified')):
            fields.append(f)
    return {'title': title, 'context': context, 'fields': fields}


def _split_numbers(s):
    """文本 -> 数值列表: "1，2，3" / "2 3.5" 拆分; 含非数值返回 []。
    手写数字 OCR 常把小数点读成横线 (37-6 -> 37.6), 数字间的
    -/—/– 按小数点容错。"""
    import re
    parts = [p.strip() for p in re.split(r'[，,、;；\s]+', str(s or ''))
             if p.strip()]
    out = []
    for p in parts:
        p = re.sub(r'(?<=\d)[-—–](?=\d)', '.', p.replace('%', ''))
        try:
            out.append(float(p) if '.' in p else int(p))
        except ValueError:
            return []
    return out


def parse_data_sheet(path):
    """数据简表 (简单表/复杂表) -> {title, context:{字段:值}, fields:[...]}。
    解析结果用完即弃, 不落盘。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    best = None
    for name in wb.sheetnames:
        ws = wb[name]
        merged = {}    # (row, col) -> 合并锚点值 (表头跨列合并时展开列角色)
        expanded = {}  # (row, col) -> 锚点(row, col), 仅原本为空的展开格
        for rng in ws.merged_cells.ranges:
            v = ws.cell(row=rng.min_row, column=rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged[(r, c)] = v
                    if (r, c) != (rng.min_row, rng.min_col):
                        expanded[(r, c)] = (rng.min_row, rng.min_col)
        entries = []  # (原行号, 值列表, number_format 列表)
        for row in ws.iter_rows():
            vals, frow = [], []
            for c in row:
                v = c.value
                if v is None and (c.row, c.column) in merged:
                    v = merged[(c.row, c.column)]
                vals.append(v)
                frow.append(c.number_format or '')
            if any(v is not None and str(v).strip() for v in vals):
                entries.append((row[0].row, vals, frow))
        fidx = {rn for rn, _, _ in entries}
        exp_f = {}  # 仅保留非空行的展开格, 键沿用原行号
        for (r, c), (ar, ac) in expanded.items():
            if r in fidx and ar in fidx:
                exp_f[(r, c)] = (ar, ac)
        sheet = _parse_sheet_rows(entries, exp_f)
        if sheet and (best is None
                      or len(sheet['fields']) > len(best['fields'])):
            best = sheet
    if not best:
        return {'title': None, 'context': {}, 'fields': []}
    # 时间类上下文字段的 Excel 日期序列号 -> "2026年11月11日"
    ctx = best.get('context')
    if ctx:
        import openpyxl.utils.datetime as xl_dt
        fixed = {}
        for k, v in ctx.items():
            if ('时间' in k or '日期' in k) \
                    and isinstance(v, (int, float)) and not isinstance(v, bool) \
                    and 20000 <= v <= 80000:
                d = xl_dt.from_excel(v)
                fixed[k] = f'{d.year}年{d.month}月{d.day}日'
            elif isinstance(v, str):
                fixed[k] = v.strip()
            else:
                fixed[k] = v
        best['context'] = fixed
    return best


def _parse_sheet_rows(entries, expanded=None):
    """一张简表的行值 -> {title, context, fields}; 识别不出表头返回 None。

    entries  [(原行号, 值列表, number_format 列表)], expanded
    {(行号, 列): 锚点(行号, 列)} 记录合并展开的空格 (实测列去重用)。"""
    import re

    def hnorm(s):
        return re.sub(r'\s+', '', str(s or ''))

    def num(v, fmt=''):
        if v is None:
            return None
        if isinstance(v, bool):
            return str(v)
        if isinstance(v, (int, float)):
            x = float(v)
            if '%' in fmt and abs(x) <= 1.0001:
                x *= 100
            return int(x) if x.is_integer() else round(x, 4)
        s = str(v).strip().replace('%', '')
        try:
            return float(s) if ('.' in s) else int(s)
        except ValueError:
            return str(v).strip()

    def values_of(v, fmt=''):
        """单元格 -> 值列表: "1，2，3" / "2 ,3 ,4" 这类文本拆成多个数值,
        拆不出数值则原样返回单个值。"""
        if v is None:
            return []
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return [num(v, fmt)]
        s = str(v).strip()
        if not s:
            return []
        parts = [p.strip() for p in re.split(r'[，,、;；\s]+', s) if p.strip()]
        nums = []
        ok = bool(parts)
        for p in parts:
            try:
                nums.append(float(p) if ('.' in p) else int(p))
            except ValueError:
                ok = False
                break
        if ok and nums:
            return nums
        n = num(v, fmt)
        return [n] if n is not None else []

    title = None
    header_idx = None
    roles = None
    title_kw = ('记录表', '检验表', '检测报告', '评定表', '验收表', '检查表')
    label_kw = ('检测项目', '检查项目', '项目名称', '项目', '名称', '指标')
    ctx_kw = ('单位', '合同段', '人', '时间', '部位', '编号', '桩号',
              '日期', '名称', '签字', '意见')
    for i, (rn, vals, fmtr) in enumerate(entries):
        norms = [hnorm(v) for v in vals]
        if header_idx is None:
            for v in norms:
                if len(v) >= 6 and any(k in v for k in title_kw):
                    if title is None or len(v) > len(title):
                        title = v
            li = next((j for j, v in enumerate(norms)
                       if v and any(k in v for k in label_kw)), None)
            si = [j for j, v in enumerate(norms)
                  if '规定值' in v or '标准值' in v or '允许偏差' in v]
            mi = [j for j, v in enumerate(norms) if '实测' in v]
            pi = [j for j, v in enumerate(norms)
                  if '合格率' in v or '通过率' in v]
            if li is not None and (si or mi):
                header_idx = i
                roles = {'label': li, 'specified': si[0] if si else None,
                         'measured': mi, 'pass': pi[0] if pi else None}

    if header_idx is None:
        return None

    # 上下文区: label格 (短文本, 命中关键词, 无数字) -> 同排右侧最近非空格
    context = {}
    for rn, vals, fmtr in entries[:header_idx]:
        norms = [hnorm(v) for v in vals]
        j = 0
        while j < len(vals):
            n = norms[j]
            if not n:
                j += 1
                continue
            if ('：' in n or ':' in n):  # 单格 "编号：JB12-4.5.2-4"
                head = re.split(r'[：:]', n)[0]
                if 1 <= len(head) <= 10 and any(k in head for k in ctx_kw):
                    context[head] = re.split(r'[：:]', n, 1)[1] or None
                    j += 1
                    continue
            if 1 <= len(n) <= 15 and any(k in n for k in ctx_kw) \
                    and not re.search(r'\d', n):
                k = j + 1
                while k < len(vals) and not norms[k]:
                    k += 1
                if k < len(vals):
                    context[n] = vals[k]
                    j = k + 1
                    continue
            j += 1

    fields = []
    for rn, vals, fmtr in entries[header_idx + 1:]:
        if roles['label'] >= len(vals):
            continue
        name = vals[roles['label']]
        if name is None or not str(name).strip():
            continue
        f = {'name': str(name).strip()}
        if roles['specified'] is not None and roles['specified'] < len(vals) \
                and vals[roles['specified']] is not None:
            f['specified'] = num(vals[roles['specified']],
                                 fmtr[roles['specified']])
        measured = []
        for j in roles['measured']:
            if j >= len(vals):
                continue
            # 合并展开的重复列 (表头跨列合并把同一实测格映到多列;
            # expanded 键是 1-based 列号, j 是 0-based 下标)
            anc = expanded.get((rn, j + 1)) if expanded else None
            if anc and (anc[1] - 1) in roles['measured'] and anc[1] - 1 != j:
                continue
            measured.extend(values_of(vals[j], fmtr[j] if j < len(fmtr) else ''))
        if measured:
            f['measured'] = measured
        if roles['pass'] is not None and roles['pass'] < len(vals) \
                and vals[roles['pass']] is not None:
            f['pass_rate'] = num(vals[roles['pass']],
                                 fmtr[roles['pass']]
                                 if roles['pass'] < len(fmtr) else '')
        fields.append(f)
    return {'title': title, 'context': context, 'fields': fields}
