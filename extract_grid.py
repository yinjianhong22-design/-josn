# -*- coding: utf-8 -*-
"""
extract_grid.py

图片/扫描件: OpenCV 提取横竖直线 -> 交点 -> 小格子(含合并单元格), 版面分析
(文字块/行切分)约束字号与位置, OCR 识别文字。
电子档 PDF: 文字层 span(字体/字号/颜色/基线) + 矢量线段直接提取, 免 OCR。
Word (.docx): 段落遍历 + 样式继承链(run->段落样式->文档默认), 标题层级按
样式名判定, 表格处理 gridSpan/vMerge 合并。
Excel (.xlsx): 逐 sheet 遍历, 值/公式/样式/合并区域分开记录。
所有格式统一组织成标准 JSON (data.json), 由 JSON 还原整页 PDF 和 HTML。

可回填表单模板 (docx/xlsx): Office 通路上用文本锚点把探测网格对齐到
原生 workbook, 给格子赋稳定地址 addr (Sheet1!B2) 和 value/formula/
number_format; 再做角色标注 (title|header|label|value|data) 与字段配对,
产出 fields (标签->填空格/内联冒号) 和 table_fields (表体行->可填列)。
data['form'] 是可索引视图: 每条字段一个扁平文档 (唯一 id/name/mode/addr/
几何平铺, 同名带 context 消歧); data['index'] 是可填充索引, 表头行为
分界线: header_fields (以上, 目录 context 纯字典填充) + rows (以下, 数
据简表按 行标签+列名 填充) + by_label (全表兜底), uuid 确定性生成,
cells 表带几何与填充 mode (overwrite/replace/append/inline), 边界情况
全部进 report。渲染层 cells/texts/rules 保持不变。
回填契约: 按 addr 找到格子, 改 text、删 lines, 几何与边框不动。

用法:
    python extract_grid.py input.png [-o out_dir] [--baidu-ak ... --baidu-sk ...]
    python extract_grid.py input.pdf [--page 0] [--dpi 300]
    python extract_grid.py input.docx | input.xlsx

依赖:
    pip install opencv-python numpy rapidocr-onnxruntime
    (PDF 输入另需: pip install pymupdf)

输出 (输出目录下):
    lines.json   横线/竖线坐标、交点、每个格子(含 rowspan/colspan/文字)
    overlay.png  在原图上画出线、交点、格子编号，便于人工核对
    data.json    标准 JSON: dpi + 整页几何 + 格子(边框/文字) + 格外游离文字
    rebuilt.pdf  按 data.json 重画的矢量 PDF 整页 (表格 + 页眉标题页码)
    data.js      同一份数据包成 window.TABLE_DATA, 供 file:// 直接打开
    index.html   纯渲染层: 读 data.js, 按 JSON 映射还原出像素级表格
"""

import argparse
import json
import os
import sys

import cv2
import numpy as np


# ---------------------------------------------------------------------------
# 加载
# ---------------------------------------------------------------------------

def _import_pymupdf():
    """导入 PyMuPDF (新版包名 pymupdf, 旧版 fitz)。"""
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        try:
            import fitz
            return fitz
        except ImportError:
            sys.exit('读取 PDF 需要 PyMuPDF:  pip install pymupdf')


def load_image(path, page=0, dpi=300):
    """加载图片；PDF 先用 PyMuPDF 按指定 dpi 渲染成位图。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pdf':
        fitz = _import_pymupdf()
        doc = fitz.open(path)
        if not (0 <= page < len(doc)):
            sys.exit(f'PDF 共 {len(doc)} 页, 无效的页码 {page}')
        pix = doc[page].get_pixmap(dpi=dpi)
        img = np.frombuffer(pix.samples, dtype=np.uint8)
        img = img.reshape(pix.height, pix.width, pix.n)
        if pix.n == 4:
            img = cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
        elif pix.n == 1:
            img = cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
        doc.close()
        return img
    img = cv2.imread(path)
    if img is None:
        sys.exit(f'无法读取图片: {path}')
    return img


# ---------------------------------------------------------------------------
# 二值化 + 横竖线提取
# ---------------------------------------------------------------------------

def binarize(gray):
    """自适应阈值二值化, 返回线条/文字为白色(255)的图。"""
    gray = cv2.GaussianBlur(gray, (3, 3), 0)
    bw = cv2.adaptiveThreshold(
        gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 15, 10)
    return bw


def extract_line_masks(bw, h_ratio=0.1, v_ratio=0.03):
    """形态学开运算分离横线和竖线。

    核长度按图像尺寸比例取: 横线核较长(10%), 因为表格横线通常贯穿表宽;
    竖线核较短(3%), 因为表格里存在只跨一两行的短竖线(表头分栏等),
    核太长会把它们当噪声滤掉。文字、噪点、短笔画会被滤掉。
    返回 (横线图, 竖线图)。
    """
    h, w = bw.shape
    h_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (max(10, int(w * h_ratio)), 1))
    v_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (1, max(10, int(h * v_ratio))))
    horiz = cv2.morphologyEx(bw, cv2.MORPH_OPEN, h_kernel)
    vert = cv2.morphologyEx(bw, cv2.MORPH_OPEN, v_kernel)
    return horiz, vert


def close_gaps(mask, axis):
    """沿线条方向闭运算, 接合同方向的断口 (如扫描件缺墨、虚线)。"""
    h, w = mask.shape
    if axis == 'h':
        kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(10, w // 50), 1))
    else:
        kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(10, h // 50)))
    return cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel)


def find_line_positions(mask, axis, min_len=20, gap=3, min_span=0.25):
    """从单方向线图中找出每条线的精确坐标。

    按像素投影找条带, 条带内按投影强度加权取中心作为线的坐标。
    min_span 过滤掉文字行这类"只覆盖一小段长度"的假线。
    返回升序坐标列表。
    """
    proj = mask.sum(axis=1 if axis == 'h' else 0) // 255  # 每行/列上的线像素数

    # 投影阈值取最大值的 4%: 表里存在只跨一两行的短线(表头分栏),
    # 阈值按最大长度比例取太高会把它们整批滤掉; 下限 min_len 防噪点
    thresh = max(min_len, proj.max() * 0.04) if proj.max() > 0 else min_len
    on = proj >= thresh
    positions = []
    i, n = 0, len(on)
    while i < n:
        if on[i]:
            j = i
            while j + 1 < n and (on[j + 1] or
                                 (j + gap < n and on[j + 1:j + gap + 1].any())):
                j += 1
            # 覆盖度: 条带内沿线条延伸方向有多大比例的列/行有线像素
            if axis == 'h':
                span = mask[i:j + 1, :].any(axis=0).mean()
            else:
                span = mask[:, i:j + 1].any(axis=1).mean()
            if span >= min_span:
                band = proj[i:j + 1].astype(np.float64)
                center = np.average(np.arange(i, j + 1), weights=band)
                positions.append(int(round(center)))
            i = j + 1
        else:
            i += 1
    return positions


def filter_dangling_lines(pos, mask, axis, other_pos, tol=25):
    """滤掉两端都不"挂靠"在另一方向线上的假线(如表格外的人名下划线)。

    取该线的实际像素延伸范围, 检查起止点 tol 像素内是否有对向线。
    真正的表格线(哪怕只跨几列)两端都会落在竖线/横线上; 下划线这类
    独立线条两端悬空, 被滤掉。
    """
    kept = []
    for p in pos:
        if axis == 'h':
            band = mask[max(0, p - 2):p + 3, :]
            cols = np.where(band.any(axis=0))[0]
        else:
            band = mask[:, max(0, p - 2):p + 3]
            cols = np.where(band.any(axis=1))[0]
        if cols.size == 0:
            continue
        lo, hi = cols[0], cols[-1]
        anchored = any(abs(o - lo) <= tol for o in other_pos) and \
                   any(abs(o - hi) <= tol for o in other_pos)
        if anchored:
            kept.append(p)
    dropped = len(pos) - len(kept)
    if dropped:
        print(f'滤掉 {dropped} 条两端悬空的假线')
    return kept


# ---------------------------------------------------------------------------
# 版面分析 (文字区域 / 图片块 / 行切分)
# ---------------------------------------------------------------------------

def analyze_layout(bw, horiz_c, vert_c, min_area=200):
    """OpenCV 版面分析: 页面内容边界 + 文字块/图片块外接矩形。

    从二值图中剔除表格线得到纯内容掩码; 文字经闭运算聚块后取轮廓。
    填充密度高的块判为图片块, 其余为文字块。
    返回 (纯内容掩码, 块列表, 页面内容边界)。
    """
    content = bw.copy()
    content[horiz_c > 0] = 0
    content[vert_c > 0] = 0
    h, w = content.shape
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT,
                                       (max(15, w // 80), max(8, h // 150)))
    merged = cv2.morphologyEx(content, cv2.MORPH_CLOSE, kernel)
    contours, _ = cv2.findContours(merged, cv2.RETR_EXTERNAL,
                                   cv2.CHAIN_APPROX_SIMPLE)
    blocks = []
    for c in contours:
        x, y, cw, ch = cv2.boundingRect(c)
        if cw * ch < min_area:
            continue
        density = content[y:y + ch, x:x + cw].mean() / 255
        blocks.append({'x': int(x), 'y': int(y), 'w': int(cw), 'h': int(ch),
                       'type': 'image' if density > 0.5 else 'text'})
    blocks.sort(key=lambda b: (b['y'], b['x']))

    cols = np.where(content.any(axis=0))[0]
    rows = np.where(content.any(axis=1))[0]
    bounds = None
    if cols.size and rows.size:
        bounds = {'x': int(cols[0]), 'y': int(rows[0]),
                  'w': int(cols[-1] - cols[0]), 'h': int(rows[-1] - rows[0])}
    return content, blocks, bounds


def split_text_lines(content, rect, gap=2, min_h=5):
    """横向投影(波峰波谷)在矩形内切出每行文字, 返回整页坐标 [(y0, y1)]。

    行高(y1-y0)可直接用于约束还原文本的字号: size_pt ≈ 行高*72/dpi*0.75。
    """
    x, y, w, h = rect
    sub = content[y:y + h, x:x + w]
    if sub.size == 0:
        return []
    on = (sub > 0).sum(axis=1) > 0
    runs = []
    i, n = 0, len(on)
    while i < n:
        if on[i]:
            j = i
            while j + 1 < n and (on[j + 1] or
                                 (j + gap < n and on[j + 1:j + gap + 1].any())):
                j += 1
            if j - i + 1 >= min_h:
                runs.append((y + i, y + j + 1))
            i = j + 1
        else:
            i += 1
    return runs


# ---------------------------------------------------------------------------
# 交点与格子
# ---------------------------------------------------------------------------

def compute_intersections(h_pos, v_pos, horiz, vert, win=5):
    """计算所有横竖线交点。若交点邻域内横竖两个方向都有线像素则保留。"""
    points = []
    for y in h_pos:
        for x in v_pos:
            y0, y1 = max(0, y - win), min(horiz.shape[0], y + win + 1)
            x0, x1 = max(0, x - win), min(horiz.shape[1], x + win + 1)
            if horiz[y0:y1, x0:x1].any() and vert[y0:y1, x0:x1].any():
                points.append({'x': x, 'y': y})
    return points


def _edge_present(mask, p1, p2, axis, thickness=2, density=0.4):
    """检查 p1->p2 之间是否真的有线段(采样投影密度)。"""
    if axis == 'h':
        y = (p1[1] + p2[1]) // 2
        band = mask[max(0, y - thickness):y + thickness + 1,
                    min(p1[0], p2[0]):max(p1[0], p2[0]) + 1]
    else:
        x = (p1[0] + p2[0]) // 2
        band = mask[min(p1[1], p2[1]):max(p1[1], p2[1]) + 1,
                    max(0, x - thickness):x + thickness + 1]
    if band.size == 0:
        return False
    proj = band.any(axis=0 if axis == 'h' else 1)
    return proj.mean() >= density


def find_cells(h_pos, v_pos, horiz, vert):
    """推出小格子(含合并单元格)。

    把相邻横竖线围成的原子块按"公共边是否缺线"做并查集合并:
    没有分隔线的相邻块属于同一个合并单元格。
    只保留至少含一个">=2 条边"的块的组, 表格外的开放区域因此被丢弃。
    """
    R, C = len(h_pos) - 1, len(v_pos) - 1
    if R <= 0 or C <= 0:
        return []

    def tid(i, j):
        return i * C + j

    parent = list(range(R * C))

    def find(a):
        while parent[a] != a:
            parent[a] = parent[parent[a]]
            a = parent[a]
        return a

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    def edges(i, j):
        tl = (v_pos[j], h_pos[i])
        tr = (v_pos[j + 1], h_pos[i])
        bl = (v_pos[j], h_pos[i + 1])
        br = (v_pos[j + 1], h_pos[i + 1])
        return (
            _edge_present(horiz, tl, tr, 'h'),  # top
            _edge_present(horiz, bl, br, 'h'),  # bottom
            _edge_present(vert, tl, bl, 'v'),   # left
            _edge_present(vert, tr, br, 'v'),   # right
        )

    tiles = {}
    for i in range(R):
        for j in range(C):
            e = edges(i, j)
            tiles[tid(i, j)] = e
            if j + 1 < C and not e[3]:      # 右边没分隔线 -> 与右邻同属一格
                union(tid(i, j), tid(i, j + 1))
            if i + 1 < R and not e[1]:      # 下边没分隔线 -> 与下邻同属一格
                union(tid(i, j), tid(i + 1, j))

    groups = {}
    for i in range(R):
        for j in range(C):
            groups.setdefault(find(tid(i, j)), []).append((i, j))

    cells = []
    for members in groups.values():
        if max(sum(tiles[tid(i, j)]) for i, j in members) < 2:
            continue  # 表格外的开放区域
        is_ = [m[0] for m in members]
        js = [m[1] for m in members]
        i0, i1, j0, j1 = min(is_), max(is_), min(js), max(js)
        x, y = v_pos[j0], h_pos[i0]
        cells.append({
            'row': i0, 'col': j0,
            'rowspan': i1 - i0 + 1, 'colspan': j1 - j0 + 1,
            'x': x, 'y': y,
            'w': v_pos[j1 + 1] - x, 'h': h_pos[i1 + 1] - y,
        })
    cells.sort(key=lambda c: (c['row'], c['col']))
    return cells


# ---------------------------------------------------------------------------
# JSON -> PDF 还原
# ---------------------------------------------------------------------------

def _find_cjk_font():
    """在常见路径里找一个支持中文的 TTF 字体 (PDF 嵌入用)。"""
    candidates = [
        r'C:\Windows\Fonts\simhei.ttf',
        r'C:\Windows\Fonts\msyh.ttc',
        r'C:\Windows\Fonts\simsun.ttc',
        '/System/Library/Fonts/PingFang.ttc',
        '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
    ]
    return next((p for p in candidates if os.path.exists(p)), None)


def _map_font(font_name):
    """把 PDF 里的原字体名映射到本机字体文件 (宋体/黑体/常见西文)。

    找不到对应文件时返回 None, 由调用方退回默认中文字体。
    """
    name = (font_name or '').lower()
    table = [
        (('simsun', 'songti', 'song'), [r'C:\Windows\Fonts\simsun.ttc']),
        (('simhei', 'heiti', 'hei'), [r'C:\Windows\Fonts\simhei.ttf']),
        (('msyh', 'yahei'), [r'C:\Windows\Fonts\msyh.ttc']),
        (('kaiti', 'kai'), [r'C:\Windows\Fonts\simkai.ttf']),
        (('fangsong', 'fang'), [r'C:\Windows\Fonts\simfang.ttf']),
        (('arial',), [r'C:\Windows\Fonts\arial.ttf']),
        (('times',), [r'C:\Windows\Fonts\times.ttf']),
    ]
    for keys, paths in table:
        if any(k in name for k in keys):
            hit = next((p for p in paths if os.path.exists(p)), None)
            if hit:
                return hit
    return None


def json_to_pdf(data, out_path, border_width=0.7):
    """按标准 JSON 重画出 PDF: 逐格画真实存在的边框, 文字自动适配填入。

    坐标换算: point = px * 72 / dpi (dpi 取自 JSON, 即渲染原图时的 dpi)。
    字体按 JSON 里记录的原字体名映射到本机字体文件 (宋体/黑体/西文),
    映射不到时退回默认中文字体; 都没有才用 PyMuPDF 内置 "china-s"
    (内置字体是逐 glyph 绘制的 fallback, 字符间会出现多余空隙)。
    文字落位三档: 带基线+字号的行 (PDF 文字层) 原样落位 > 带行高的行
    (OCR 行切分) 按行高定字号 > 无位置信息时 insert_textbox 自动适配。
    """
    fitz = _import_pymupdf()
    s = 72.0 / data.get('dpi', 300)
    doc = fitz.open()
    page = doc.new_page(width=data['width'] * s, height=data['height'] * s)
    default_file = _find_cjk_font()
    registry = {}  # fontfile -> (pdf 内字体名, fitz.Font)

    def font_for(orig_name=None):
        """按原字体名取 (pdf 内字体名, Font), 按需嵌入并缓存。"""
        file = _map_font(orig_name) or default_file
        if file:
            if file not in registry:
                fn = f'F{len(registry)}'
                page.insert_font(fontname=fn, fontfile=file)
                registry[file] = (fn, fitz.Font(fontfile=file))
            return registry[file]
        if 'china-s' not in registry:
            registry['china-s'] = ('china-s', fitz.Font('china-s'))
        return registry['china-s']

    # 电子档有原始矢量线段: 直接按精确坐标+原线宽逐条画, 每根线只画一次,
    # 不再按格子边逐边描绘 (避免共享边重复描、粗外框被画成细线、下划线丢失)
    rules = data.get('rules') or []
    for rl in rules:
        page.draw_line((rl['x0'] * s, rl['y0'] * s),
                       (rl['x1'] * s, rl['y1'] * s),
                       width=max(0.1, rl['width'] * s))

    for c in data['cells']:
        x, y = c['x'] * s, c['y'] * s
        r = fitz.Rect(x, y, x + c['w'] * s, y + c['h'] * s)
        b = c.get('borders') or {}
        if not rules:
            if b['top']:
                page.draw_line(r.tl, r.tr, width=border_width)
            if b['bottom']:
                page.draw_line(r.bl, r.br, width=border_width)
            if b['left']:
                page.draw_line(r.tl, r.bl, width=border_width)
            if b['right']:
                page.draw_line(r.tr, r.br, width=border_width)

        max_w = c['w'] * s - 4
        if c.get('lines') and 'size' in c['lines'][0]:
            # PDF 文字层: 每行按基线坐标 + 原字号/字体/颜色精确落位, 不重排
            for ln in c['lines']:
                txt = ln['text'].strip()
                if not txt:
                    continue
                fn, fnt = font_for(ln.get('font'))
                col = ln.get('color') or '#000000'
                rgb = tuple(int(col[i:i + 2], 16) / 255 for i in (1, 3, 5))
                size = ln['size']
                tw = fnt.text_length(txt, fontsize=size)
                orig_w = (ln.get('w') or 0) * s
                pt = (ln['x'] * s, ln['y'] * s)
                if len(txt) > 1 and orig_w > tw * 1.05:
                    # 原件用字符间距撑开的行 (如 "年 月 日"): 逐字均摊间距
                    gap = (orig_w - tw) / (len(txt) - 1)
                    cx = pt[0]
                    for ch in txt:
                        page.insert_text((cx, pt[1]), ch, fontsize=size,
                                         fontname=fn, color=rgb)
                        cx += fnt.text_length(ch, fontsize=size) + gap
                    continue
                if tw > max_w:  # 嵌入字体度量差异超宽时同比缩小
                    size *= max_w / tw
                page.insert_text(pt, txt, fontsize=size,
                                 fontname=fn, color=rgb)
                if ln.get('bold'):
                    page.insert_text((pt[0] + 0.3, pt[1]), txt,
                                     fontsize=size, fontname=fn, color=rgb)
            continue

        fontname, font = font_for(c.get('font'))
        if c.get('lines'):
            # OCR 行切分: 逐行按行高定字号、按行位置落基线。
            # 同一格内 >=3 行或有任一行接近格宽 -> 按段落处理(整格左对齐),
            # 避免段落里较短的行被居中造成参差不齐
            para = len(c['lines']) >= 3 or any(
                font.text_length(ln['text'].strip(),
                                 fontsize=ln['h'] * s * 0.75) > max_w * 0.7
                for ln in c['lines'] if ln['text'].strip())
            for ln in c['lines']:
                txt = ln['text'].strip()
                if not txt:
                    continue
                size = ln['h'] * s * 0.75
                tw = font.text_length(txt, fontsize=size)
                if tw > max_w:  # 字体度量差异导致超宽时同比缩小
                    size *= max_w / tw
                    tw = max_w
                lx = r.x0 + 2 if para else r.x0 + (c['w'] * s - tw) / 2
                ly = (ln['y'] + ln['h'] * 0.82) * s
                page.insert_text((lx, ly), txt, fontsize=size,
                                 fontname=fontname)
            continue

        text = (c.get('text') or '').strip()
        if not text:
            continue
        inner = fitz.Rect(r.x0 + 2, r.y0 + 1, r.x1 - 2, r.y1 - 1)
        lines = text.split('\n')
        size = min(float(c.get('size') or 12.0), inner.height * 0.8)
        if len(lines) == 1:
            # 单行文字先按宽度预缩, 避免 insert_textbox 把短行折成两行
            tw = font.text_length(text, fontsize=size)
            if tw > inner.width:
                size = max(4.0, size * inner.width / tw)
        col = c.get('color') or '#000000'
        crgb = tuple(int(col[i:i + 2], 16) / 255 for i in (1, 3, 5))
        while size > 4:
            widest = max(font.text_length(l, fontsize=size) for l in lines)
            # 显式对齐优先; 否则段落(最宽行接近格宽)左对齐, 短文字居中
            al = c.get('align')
            if al == 'left':
                align = fitz.TEXT_ALIGN_LEFT
            elif al == 'right':
                align = fitz.TEXT_ALIGN_RIGHT
            elif al == 'center':
                align = fitz.TEXT_ALIGN_CENTER
            else:
                align = fitz.TEXT_ALIGN_LEFT if widest > inner.width * 0.7 \
                    else fitz.TEXT_ALIGN_CENTER
            # 估算换行后的行数, 把文字块整体垂直居中
            n = sum(max(1, int(np.ceil(
                font.text_length(l, fontsize=size) / inner.width)))
                for l in lines)
            block_h = n * size * 1.3
            box = fitz.Rect(inner.x0, inner.y0, inner.x1, inner.y1)
            if block_h < inner.height:
                box.y0 += (inner.height - block_h) / 2
            rc = page.insert_textbox(box, text, fontsize=size,
                                     fontname=fontname, align=align,
                                     color=crgb)
            if rc >= 0:
                if c.get('bold'):  # 模拟加粗
                    page.insert_textbox(
                        fitz.Rect(box.x0 + 0.3, box.y0, box.x1 + 0.3,
                                  box.y1),
                        text, fontsize=size, fontname=fontname,
                        align=align, color=crgb)
                break
            size -= 0.5

    # 表格线框外的游离文字(页眉/标题/页码等): 按基线坐标、原字号、
    # 原字体、原色落位; 粗体/标题用微偏移二次绘制模拟加粗
    for t in data.get('texts', []):
        fn, fnt = font_for(t.get('font'))
        col = t.get('color') or '#000000'
        rgb = tuple(int(col[i:i + 2], 16) / 255 for i in (1, 3, 5))
        size = t.get('size', 10)
        tx = t['x'] * s
        tw = fnt.text_length(t['text'], fontsize=size)
        if t.get('align') and t.get('w'):  # 居中/右对齐的段落
            if t['align'] == 'center':
                tx += (t['w'] * s - tw) / 2
            elif t['align'] == 'right':
                tx += t['w'] * s - tw
        pt = (tx, t['y'] * s)
        page.insert_text(pt, t['text'], fontsize=size,
                         fontname=fn, color=rgb)
        if t.get('bold') or t.get('role') in ('title', 'heading'):
            page.insert_text((pt[0] + 0.3, pt[1]), t['text'], fontsize=size,
                             fontname=fn, color=rgb)

    err = None
    for k in range(10):
        # 目标文件被占用(如在 PDF 阅读器中打开)时自动换名: _1, _2, ...
        name = out_path if k == 0 else out_path.replace('.pdf', f'_{k}.pdf')
        try:
            doc.save(name, deflate=True)
            if k:
                print(f'{out_path} 被占用, 已改存 {name}')
            out_path = name
            break
        except Exception as e:
            err = e
    else:
        doc.close()
        raise err
    doc.close()
    return out_path


# ---------------------------------------------------------------------------
# PDF 文字层提取 (免 OCR, 坐标精准)
# ---------------------------------------------------------------------------

def _classify_line_roles(lines, page_w):
    """字号聚类 + 坐标规则, 给每行文字定语义角色 (title/heading/body)。

    以字号中位数为正文基准: 明显大于正文且接近页面水平居中 -> 标题;
    明显大于正文但不居中 -> 小标题; 其余为正文。
    """
    if not lines:
        return
    sizes = sorted(l['size'] for l in lines)
    body = sizes[len(sizes) // 2]
    for l in lines:
        centered = abs(l['cx'] - page_w / 2) < page_w * 0.12
        if l['size'] >= body * 1.3:
            l['role'] = 'title' if centered else 'heading'
        elif l['size'] >= body * 1.12:
            l['role'] = 'heading'
        else:
            l['role'] = 'body'


def extract_pdf_text(path, page, dpi, cells):
    """探测并读取 PDF 文字层: 电子档提取 span 级信息, 扫描件返回 None。

    每个 span 带字体/字号/颜色/加粗/基线坐标(PDF point -> 渲染像素换算,
    缩放比 = dpi/72)。以文字行中心点判定归属格子; 格内多行排序拼接。
    格外行(页眉/标题/页码等)按"字号聚类 + 坐标规则"标注语义角色后保留。
    返回 ({(row,col): 文本}, [游离文字]); 无文字层(扫描件)返回 (None, [])。
    """
    fitz = _import_pymupdf()
    doc = fitz.open(path)
    if not (0 <= page < len(doc)):
        doc.close()
        return None, []
    pg = doc[page]
    d = pg.get_text('rawdict')
    page_w = pg.rect.width * dpi / 72.0
    doc.close()

    scale = dpi / 72.0
    lines = []
    for b in d['blocks']:
        if b.get('type') != 0:
            continue
        for l in b['lines']:
            text = ''.join(ch['c'] for s in l['spans']
                           for ch in s.get('chars', [])).strip()
            if not text:
                continue
            # 逐字符定位: WPS/Excel 转出的行常用前导空格做缩进/对齐,
            # span 的 origin/bbox 把空格宽度也算进去, 直接用会让文字
            # 整体左移、宽度虚增(触发错误的逐字均摊); 改用非空格字符
            # 的实际 bbox 和首个非空格字符的基线起点
            kept = [(s, ch) for s in l['spans'] for ch in s.get('chars', [])
                    if ch['c'].strip()]
            x0 = min(ch['bbox'][0] for _, ch in kept) * scale
            y0 = min(ch['bbox'][1] for _, ch in kept) * scale
            x1 = max(ch['bbox'][2] for _, ch in kept) * scale
            y1 = max(ch['bbox'][3] for _, ch in kept) * scale
            ox, oy = kept[0][1]['origin']
            seen, spans_k = set(), []
            for s, _ in kept:
                if id(s) not in seen:
                    seen.add(id(s))
                    spans_k.append(s)
            s0 = max(spans_k, key=lambda s: s['size'])  # 主导 span
            # 分段信息: 同一行里字体/位置不同的 span 分开记录,
            # 游离文字按段独立落位(如 "监理单位:___  编号:JB12..." 一行两段)
            segs = []
            for s in spans_k:
                schars = [ch for ch in s.get('chars', []) if ch['c'].strip()]
                segs.append({
                    'text': ''.join(ch['c'] for ch in s['chars']).strip(),
                    'ox': schars[0]['origin'][0] * scale,
                    'oy': schars[0]['origin'][1] * scale,
                    'size': s['size'], 'font': s['font'],
                    'color': '#%06x' % s['color'],
                    'bold': bool(s['flags'] & 16)
                            or 'bold' in s['font'].lower(),
                })
            lines.append({
                'cx': (x0 + x1) / 2, 'cy': (y0 + y1) / 2,
                'x0': x0, 'y0': y0, 'x1': x1, 'y1': y1,
                'ox': ox * scale, 'oy': oy * scale,
                'size': s0['size'], 'font': s0['font'],
                'color': '#%06x' % s0['color'],
                'bold': bool(s0['flags'] & 16)
                        or 'bold' in s0['font'].lower(),
                'text': text, 'segs': segs,
                'vertical': abs(l['dir'][0]) < 0.5,
            })
    if not lines:
        return None, []

    _classify_line_roles(lines, page_w)

    texts = {}
    assigned = set()
    for c in cells:
        inside = []
        for k, l in enumerate(lines):
            if (c['x'] - 2 <= l['cx'] <= c['x'] + c['w'] + 2
                    and c['y'] - 2 <= l['cy'] <= c['y'] + c['h'] + 2):
                inside.append(l)
                assigned.add(k)
        if not inside:
            continue
        vertical = any(l['vertical'] for l in inside)
        # 横排按 y 再按 x 排序; 竖排按 x 再按 y 排序
        inside.sort(key=lambda l: (-l['x0'], l['y0']) if vertical
                    else (l['y0'], l['x0']))
        # 每行保留基线坐标/字号/字体/颜色/原始宽度(供还原字符间距)
        c['lines'] = [{'x': l['ox'], 'y': l['oy'], 'size': l['size'],
                       'font': l['font'], 'color': l['color'],
                       'bold': l['bold'], 'text': l['text'],
                       'w': l['x1'] - l['x0']}
                      for l in inside]
        c['font'] = max(inside, key=lambda l: l['size'])['font']  # 主导字体
        if vertical:
            texts[(c['row'], c['col'])] = '\n'.join(l['text'] for l in inside)
            continue
        # 同一视觉行(y 接近)的多段文字横向合并 (如 "年  月  日" 是 3 个 span)
        merged = []
        for l in inside:
            tol = max(6.0, l['size'] * scale * 0.5)
            if merged and abs(l['y0'] - merged[-1][0]['y0']) <= tol:
                merged[-1].append(l)
            else:
                merged.append([l])
        rows = []
        for grp in merged:
            grp.sort(key=lambda l: l['x0'])
            rows.append(''.join(l['text'] for l in grp))
        texts[(c['row'], c['col'])] = '\n'.join(rows)

    # 游离文字按 span 分段输出: 每段带自己的基线/字体/颜色,
    # 一行里左右分开的多段文字(编号/日期等)各自精确落位
    free = []
    for k, l in enumerate(lines):
        if k in assigned:
            continue
        for seg in l['segs']:
            free.append({'x': seg['ox'], 'y': seg['oy'], 'size': seg['size'],
                         'font': seg['font'], 'color': seg['color'],
                         'bold': seg['bold'], 'role': l['role'],
                         'text': seg['text']})
    free.sort(key=lambda t: (t['y'], t['x']))
    return texts, free


def _merge_positions(a, b, tol=6):
    """合并两个线坐标列表 (OpenCV 检测值与矢量线段聚类值), 相近的取平均。"""
    vals = sorted(list(a) + list(b))
    groups = []
    for v in vals:
        if groups and v - groups[-1][-1] <= tol:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [int(round(sum(g) / len(g))) for g in groups]


def rules_to_masks(rules, shape, tol=6):
    """把矢量线段渲染成横/竖线掩码, 并直接聚出线坐标。

    用于电子档 PDF 中 OpenCV 形态学检测失效的场景 (如只占页面一角的
    小表格: 核长/覆盖度按整页比例取, 会把短线整批滤掉)。
    返回 (横线掩码, 竖线掩码, 横线坐标, 竖线坐标)。
    """
    horiz = np.zeros(shape, np.uint8)
    vert = np.zeros(shape, np.uint8)
    hs, vs = [], []
    for rl in rules:
        x0, y0, x1, y1 = rl['x0'], rl['y0'], rl['x1'], rl['y1']
        th = max(1, int(round(rl.get('width', 1))))
        if abs(y1 - y0) <= 2 and abs(x1 - x0) > 3:
            y = int(round((y0 + y1) / 2))
            cv2.line(horiz, (int(min(x0, x1)), y), (int(max(x0, x1)), y),
                     255, th)
            hs.append(y)
        elif abs(x1 - x0) <= 2 and abs(y1 - y0) > 3:
            x = int(round((x0 + x1) / 2))
            cv2.line(vert, (x, int(min(y0, y1))), (x, int(max(y0, y1))),
                     255, th)
            vs.append(x)

    def cluster(vals):
        vals.sort()
        groups = []
        for v in vals:
            if groups and v - groups[-1][-1] <= tol:
                groups[-1].append(v)
            else:
                groups.append([v])
        return [int(round(sum(g) / len(g))) for g in groups]

    return horiz, vert, cluster(hs), cluster(vs)


def extract_pdf_rules(path, page, dpi):
    """提取 PDF 里的矢量线段(表格线/下划线等), 返回整页像素坐标线段列表。

    电子档的线本来就是矢量数据, 坐标和线宽都是精确值; 只保留横平竖直的
    线段和细长矩形(填充式粗线)。返回 [{'x0','y0','x1','y1','width'}],
    坐标为渲染图像素, width 同样按 dpi/72 换算成像素。
    """
    fitz = _import_pymupdf()
    doc = fitz.open(path)
    if not (0 <= page < len(doc)):
        doc.close()
        return []
    drawings = doc[page].get_drawings()
    doc.close()

    scale = dpi / 72.0
    rules = []
    for d in drawings:
        w = max(d.get('width') or 0.5, 0.1)
        for item in d['items']:
            if item[0] == 'l':              # 线段
                p1, p2 = item[1], item[2]
                x0, y0, x1, y1 = p1.x, p1.y, p2.x, p2.y
            elif item[0] == 're':           # 细长矩形 = 填充的线
                r = item[1]
                if r.height <= 3 and r.width > 5:
                    y0 = y1 = (r.y0 + r.y1) / 2
                    x0, x1 = r.x0, r.x1
                    w = max(w, r.height)
                elif r.width <= 3 and r.height > 5:
                    x0 = x1 = (r.x0 + r.x1) / 2
                    y0, y1 = r.y0, r.y1
                    w = max(w, r.width)
                else:
                    continue
            else:
                continue
            if abs(y0 - y1) <= 1 or abs(x0 - x1) <= 1:
                rules.append({'x0': x0 * scale, 'y0': y0 * scale,
                              'x1': x1 * scale, 'y1': y1 * scale,
                              'width': w * scale})
    return rules


# ---------------------------------------------------------------------------
# OCR (RapidOCR, 支持中文印刷体与较工整的手写体)
# ---------------------------------------------------------------------------

def _tight_crop(crop, margin=8):
    """裁掉留白, 只保留内容包围盒(含少量边距)。"""
    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    dark = np.column_stack(np.where(gray < 128))
    if dark.size == 0:
        return None
    y0, x0 = dark.min(axis=0)
    y1, x1 = dark.max(axis=0)
    return crop[max(0, y0 - margin):y1 + margin,
                max(0, x0 - margin):x1 + margin]


def _prepare_crop(img, c, inset=4):
    """从原图裁出一个格子: 内缩避开边框, 空格跳过(返回 None),
    极端长宽比补白到不超过 5:1 (否则检测器会退化为整格单行识别)。
    """
    x0, y0 = c['x'] + inset, c['y'] + inset
    x1, y1 = c['x'] + c['w'] - inset, c['y'] + c['h'] - inset
    if x1 <= x0 or y1 <= y0:
        return None
    crop = img[y0:y1, x0:x1]
    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    if int((gray < 128).sum()) < 30:
        return None  # 深色像素太少, 格内基本没有内容
    ch, cw = crop.shape[:2]
    if cw > ch * 5:
        pad = (cw // 5 - ch) // 2
        crop = cv2.copyMakeBorder(crop, pad, pad, 0, 0,
                                  cv2.BORDER_CONSTANT, value=(255, 255, 255))
    elif ch > cw * 5:
        pad = (ch // 5 - cw) // 2
        crop = cv2.copyMakeBorder(crop, 0, 0, pad, pad,
                                  cv2.BORDER_CONSTANT, value=(255, 255, 255))
    return crop


def ocr_cells(img, cells, inset=4):
    """逐格 OCR (RapidOCR 本地): 按格子坐标裁出小图单独识别。

    整图一次识别容易在表格场景漏检/串行, 逐格识别区域小、干扰少。
    检测对单字格会漏检, 这种情况退化为"紧贴内容整格识别"。
    返回 {(row,col): 文本}; 未安装 rapidocr-onnxruntime 时返回 None。
    """
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        return None
    try:
        engine = RapidOCR(text_score=0.3)
    except Exception:
        engine = RapidOCR()
    texts = {}
    for c in cells:
        crop = _prepare_crop(img, c, inset)
        if crop is None:
            continue
        try:
            res, _ = engine(crop)
        except Exception:
            res = None
        if res:
            lines = sorted(res, key=lambda r: sum(p[1] for p in r[0]) / 4.0)
            texts[(c['row'], c['col'])] = '\n'.join(t for _, t, _ in lines)
            continue
        # 检测漏检 (单字/连体手写等): 紧贴内容, 整格作为一行直接识别
        sub = _tight_crop(crop)
        if sub is None:
            continue
        try:
            rec_res, _ = engine.text_recognizer([sub])
        except Exception:
            continue
        if rec_res and float(rec_res[0][1]) >= 0.3:
            texts[(c['row'], c['col'])] = rec_res[0][0]
    return texts


def ocr_cells_paddle(img, cells, url, token, inset=4):
    """逐格 OCR (PaddleOCR 官方 API / PP-OCRv5+): 每格裁图后单独请求。

    API_URL 与 token 在 https://aistudio.baidu.com/paddleocr/task 的
    "API 调用示例" 中获取。返回 {(row,col): 文本}; 网络失败返回 None。
    """
    import base64

    import requests

    headers = {'Authorization': f'token {token}',
               'Content-Type': 'application/json'}
    texts = {}
    for idx, c in enumerate(cells):
        crop = _prepare_crop(img, c, inset)
        if crop is None:
            continue
        b64 = base64.b64encode(cv2.imencode('.png', crop)[1]).decode('ascii')
        payload = {
            'file': b64, 'fileType': 1, 'visualize': False,
            'useDocOrientationClassify': False,
            'useDocUnwarping': False,
            'useTextlineOrientation': False,
        }
        try:
            r = requests.post(url, json=payload, headers=headers, timeout=120)
            r.raise_for_status()
            body = r.json()
        except Exception as e:
            print(f'  格[{c["row"]},{c["col"]}] OCR 请求失败: {e}')
            continue
        if body.get('errorCode', 0) != 0:
            print(f'  格[{c["row"]},{c["col"]}] OCR 报错: {body.get("errorMsg")}')
            continue
        results = body.get('result', {}).get('ocrResults') or []
        if not results:
            continue
        pr = results[0].get('prunedResult', {})
        rec_texts = pr.get('rec_texts') or []
        polys = pr.get('rec_polys') or pr.get('dt_polys') or []
        # 按文本框纵向位置排序后拼接
        items = []
        for k, t in enumerate(rec_texts):
            y = float(np.mean([p[1] for p in polys[k]])) if k < len(polys) \
                else float(k)
            items.append((y, t))
        if items:
            texts[(c['row'], c['col'])] = '\n'.join(
                t for _, t in sorted(items))
        print(f'  格[{c["row"]},{c["col"]}] ({idx + 1}/{len(cells)}) 完成')
    return texts


def _baidu_token(ak, sk):
    """用 AK/SK 换 access_token (30 天有效), 缓存到本地文件避免重复请求。"""
    import time

    import requests

    cache_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              '.baidu_token_cache.json')
    if os.path.exists(cache_path):
        try:
            with open(cache_path, encoding='utf-8') as f:
                cache = json.load(f)
            if cache.get('ak') == ak and cache.get('expires', 0) > time.time():
                return cache['token']
        except Exception:
            pass
    r = requests.get('https://aip.baidubce.com/oauth/2.0/token',
                     params={'grant_type': 'client_credentials',
                             'client_id': ak, 'client_secret': sk},
                     timeout=30)
    d = r.json()
    token = d.get('access_token')
    if not token:
        raise RuntimeError(f'百度 access_token 获取失败: {d}')
    with open(cache_path, 'w', encoding='utf-8') as f:
        # 官方有效期 30 天, 缓存提前 1 天失效
        json.dump({'ak': ak, 'token': token,
                   'expires': time.time() + d.get('expires_in', 2592000) - 86400},
                  f)
    return token


def ocr_cells_baidu(img, cells, ak, sk, api='accurate_basic', inset=4):
    """逐格 OCR (百度智能云文字识别): 每格裁图后单独请求。

    AK/SK 在百度智能云文字识别控制台的应用列表中获取。
    api 可选 accurate_basic(高精度,默认)/general_basic/handwriting 等。
    返回 {(row,col): 文本}。
    """
    import base64
    import time

    import requests

    token = _baidu_token(ak, sk)
    url = f'https://aip.baidubce.com/rest/2.0/ocr/v1/{api}'
    texts = {}
    for idx, c in enumerate(cells):
        crop = _prepare_crop(img, c, inset)
        if crop is None:
            continue
        b64 = base64.b64encode(cv2.imencode('.png', crop)[1]).decode('ascii')
        body = None
        for attempt in range(3):
            try:
                r = requests.post(
                    url, params={'access_token': token}, data={'image': b64},
                    headers={'Content-Type':
                             'application/x-www-form-urlencoded'},
                    timeout=60)
                body = r.json()
            except Exception as e:
                print(f'  格[{c["row"]},{c["col"]}] OCR 请求失败: {e}')
                break
            code = body.get('error_code', 0)
            if code == 0:
                break
            if code in (17, 18, 19):      # 限流/配额错误, 等待重试
                time.sleep(1.5 * (attempt + 1))
                continue
            print(f'  格[{c["row"]},{c["col"]}] OCR 报错: '
                  f'{body.get("error_code")} {body.get("error_msg")}')
            break
        if body and body.get('error_code', 0) == 0:
            words = [w['words'] for w in body.get('words_result', [])]
            if words:
                texts[(c['row'], c['col'])] = '\n'.join(words)
        time.sleep(0.35)  # 免费额度 QPS 有限, 控制节奏
        print(f'  格[{c["row"]},{c["col"]}] ({idx + 1}/{len(cells)}) 完成')
    return texts


def _pick_ocr_backend(args):
    """按命令行参数选择 OCR 后端, 返回 (名称, (img, cells) -> dict)。"""
    if args.baidu_ak and args.baidu_sk:
        return (f'百度智能云文字识别 ({args.baidu_api})',
                lambda img, cl: ocr_cells_baidu(img, cl, args.baidu_ak,
                                                args.baidu_sk,
                                                api=args.baidu_api))
    if args.paddle_url and args.paddle_token:
        return ('PaddleOCR API',
                lambda img, cl: ocr_cells_paddle(img, cl, args.paddle_url,
                                                 args.paddle_token))
    return 'RapidOCR (本地)', ocr_cells


def _attach_line_pos(cells, texts):
    """把 OCR 出的每行文字对号到投影切分出的行位置上。

    OCR 逐格返回多行文本但不含坐标; 这里用横向投影的行切分结果
    (cell['text_lines']) 按顺序对位, 供 PDF 还原按行高定字号和位置。
    行数对不上时放弃该格的位置约束, 退回整体自动适配。
    """
    for c in cells:
        t = texts.get((c['row'], c['col']))
        spans = c.get('text_lines') or []
        if not t or len(spans) != len(t.split('\n')):
            continue
        c['lines'] = [{'y': y0, 'h': y1 - y0, 'text': txt}
                      for (y0, y1), txt in zip(spans, t.split('\n'))]


def ocr_free_blocks(img, blocks, cells, content, backend, dpi):
    """扫描件/图片路径: OCR 表格线框外的文字块, 转成游离文字列表。

    每个块按投影切行, 行高换算字号(size_pt = 行高*72/dpi*0.75),
    基线取行底上 80% 处, 产出结构与 PDF 文字层路径的游离文字一致。
    """
    pseudo = []
    for k, b in enumerate(blocks):
        if b['type'] != 'text':
            continue
        cx, cy = b['x'] + b['w'] / 2, b['y'] + b['h'] / 2
        if any(c['x'] - 2 <= cx <= c['x'] + c['w'] + 2
               and c['y'] - 2 <= cy <= c['y'] + c['h'] + 2 for c in cells):
            continue
        pseudo.append({'row': -1, 'col': k, 'x': b['x'], 'y': b['y'],
                       'w': b['w'], 'h': b['h']})
    if not pseudo:
        return []
    print(f'格外文字块 {len(pseudo)} 个, OCR 中...')
    res = backend(img, pseudo) or {}
    free = []
    for p in pseudo:
        t = res.get((p['row'], p['col']))
        if not t:
            continue
        ols = t.split('\n')
        spans = split_text_lines(content, (p['x'], p['y'], p['w'], p['h']))
        if len(spans) != len(ols):
            step = p['h'] / len(ols)  # 行数对不上: 块高均分兜底
            spans = [(p['y'] + int(i * step), p['y'] + int((i + 1) * step))
                     for i in range(len(ols))]
        for (y0, y1), txt in zip(spans, ols):
            lh = y1 - y0
            free.append({'x': p['x'], 'y': y0 + lh * 0.8,
                         'size': lh * 72 / dpi * 0.75,
                         'color': '#000000', 'bold': False, 'role': 'body',
                         'text': txt, 'cx': p['x'] + p['w'] / 2})
    free.sort(key=lambda t: (t['y'], t['x']))
    return free


# ---------------------------------------------------------------------------
# 标准 JSON (data.json) —— 表格还原的唯一数据源
# ---------------------------------------------------------------------------

def build_table_data(cells, texts, horiz_c, vert_c, dpi, img_size,
                     free_texts=None, rules=None):
    """把识别结果组织成一份"可还原整页 PDF"的标准 JSON 数据结构。

    坐标为整页渲染图像素坐标(含页边距), dpi 字段记录了像素->PDF point
    的换算依据 (point = px * 72 / dpi)。包含:
      cells  每个格子的几何、rowspan/colspan、四周边框有无、文字
      texts  表格线框外的游离文字(页眉/标题/页码等), 带基线坐标和字号(pt)
      rules  电子档 PDF 的原始矢量线段(精确坐标+线宽); 有 rules 时渲染层
             直接画 rules, 不再按格子边逐边描绘(避免重复/偏移)
    """
    if not cells:
        return None

    def edges(c):
        tl = (c['x'], c['y'])
        tr = (c['x'] + c['w'], c['y'])
        bl = (c['x'], c['y'] + c['h'])
        br = (c['x'] + c['w'], c['y'] + c['h'])
        return {
            'top': bool(_edge_present(horiz_c, tl, tr, 'h')),
            'bottom': bool(_edge_present(horiz_c, bl, br, 'h')),
            'left': bool(_edge_present(vert_c, tl, bl, 'v')),
            'right': bool(_edge_present(vert_c, tr, br, 'v')),
        }

    w_img, h_img = img_size
    cell_list = []
    for c in cells:
        d = {
            'row': c['row'], 'col': c['col'],
            'rowspan': c['rowspan'], 'colspan': c['colspan'],
            'x': c['x'], 'y': c['y'],
            'w': c['w'], 'h': c['h'],
            'borders': edges(c),
            'text': texts.get((c['row'], c['col']), ''),
        }
        if c.get('lines'):
            # 行级信息: PDF 文字层为 基线+字号+字体+颜色, OCR 为 行高+文本
            d['lines'] = c['lines']
        if c.get('font'):
            d['font'] = c['font']  # 主导字体名, 供还原时映射本机字体
        cell_list.append(d)

    return {
        'dpi': dpi,
        'width': w_img,
        'height': h_img,
        'cells': cell_list,
        'texts': free_texts or [],
        'rules': rules or [],
    }


# ---------------------------------------------------------------------------
# Word/Excel -> PDF (本机 WPS/Office COM 排版引擎, 高保真)
# ---------------------------------------------------------------------------

def _find_wpscli():
    """在本机常见路径找 WPS 命令行工具 wpscli.exe。"""
    import glob
    patterns = [
        r'C:\Program Files*\WPS Office\*\clitool\wpscli.exe',
        r'C:\Program Files*\WPS Office\*\office6\wpscli.exe',
        os.path.expandvars(r'%LOCALAPPDATA%\Kingsoft\WPS Office\*'
                           r'\clitool\wpscli.exe'),
        os.path.expandvars(r'%LOCALAPPDATA%\Kingsoft\WPS Office\*'
                           r'\office6\wpscli.exe'),
    ]
    for pat in patterns:
        hit = glob.glob(pat)
        if hit:
            return sorted(hit)[-1]  # 取最新版本目录
    return None


def _set_attr(tag, name, value):
    """在 XML 开标签里设置/替换一个属性, 返回新标签文本。"""
    import re
    if re.search(rf'\s{name}="[^"]*"', tag):
        return re.sub(rf'\s{name}="[^"]*"', f' {name}="{value}"', tag)
    if tag.endswith('/>'):
        return tag[:-2] + f' {name}="{value}"/>'
    return tag[:-1] + f' {name}="{value}">'


def _fit_sheet_xml(data):
    """在一个 worksheet XML 里打开"调整为 1 页宽" (fitToPage)。

    只做三处原子修改: sheetPr 里加 <pageSetUpPr fitToPage="1"/>;
    pageSetup 去掉 scale、写上 fitToWidth="1" fitToHeight="0"
    (高度 0 = 纵向不限制页数)。其余字节原样保留。
    """
    import re
    xml = data.decode('utf-8')

    # 1) sheetPr/pageSetUpPr: fitToPage 开关
    if '<pageSetUpPr' in xml:
        xml = re.sub(r'<pageSetUpPr[^>]*/>',
                     lambda m: _set_attr(m.group(0), 'fitToPage', '1'), xml)
    elif re.search(r'<sheetPr[^>]*/>', xml):
        xml = re.sub(r'<sheetPr([^>]*)/>',
                     r'<sheetPr\1><pageSetUpPr fitToPage="1"/></sheetPr>', xml,
                     count=1)
    elif re.search(r'<sheetPr[^>]*>', xml):
        xml = re.sub(r'(</sheetPr>)',
                     r'<pageSetUpPr fitToPage="1"/>\1', xml, count=1)
    else:
        xml = re.sub(r'(<worksheet[^>]*>)',
                     r'\1<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>',
                     xml, count=1)

    # 2) pageSetup: 删 scale, 写 fitToWidth/fitToHeight
    def fix_page_setup(m):
        tag = re.sub(r'\sscale="[^"]*"', '', m.group(0))
        tag = _set_attr(tag, 'fitToWidth', '1')
        return _set_attr(tag, 'fitToHeight', '0')

    if '<pageSetup' in xml:
        xml = re.sub(r'<pageSetup[^>]*/>', fix_page_setup, xml)
    elif '<pageMargins' in xml:
        xml = re.sub(r'(<pageMargins[^>]*/>)',
                     r'\1<pageSetup fitToWidth="1" fitToHeight="0"/>',
                     xml, count=1)
    return xml.encode('utf-8')


def _xlsx_fit_to_page(src, dst):
    """复制 xlsx 并打开"调整为 1 页宽", 供排版引擎转 PDF 时等比缩放。

    与内容级缩放(改列宽/字号)不同: 这里只在 worksheet XML 里注入
    fitToPage 页面设置, 其余内容逐字节保持原样, 副本与原文件完全一致;
    缩放由 WPS/Excel 的排版引擎在转 PDF 时统一等比完成, 内容的相对
    位置严格不变, 生成的 PDF 就是缩放好的完整页面 (横向永不溢出分页)。
    """
    import re
    import zipfile

    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', item.filename):
                data = _fit_sheet_xml(data)
            zout.writestr(item, data)
    return dst


def office_to_pdf(path, out_dir):
    """把 docx/xlsx 转成 PDF, 排版由真正的排版引擎完成(高保真),
    转出的 PDF 走电子档通路。

    优先用 WPS 官方命令行 wpscli (word2pdf/excel2pdf, 渲染正确);
    找不到再退到 WPS/Office COM 接口; 都不可用返回 None,
    由调用方回退到 python-docx/openpyxl 简化排版。
    """
    import subprocess

    ext = os.path.splitext(path)[1].lower()
    abs_in = os.path.abspath(path)
    abs_out = os.path.abspath(os.path.join(
        out_dir, os.path.splitext(os.path.basename(path))[0] + '_office.pdf'))

    cur_in = abs_in
    if ext == '.xlsx':
        # 转换前注入"调整为 1 页宽"页面设置 (XML 级, 内容逐字节不动),
        # 转出的 PDF 就是排版引擎等比缩放好的完整页面, 横向永不溢出
        try:
            cur_in = _xlsx_fit_to_page(abs_in, abs_out.replace(
                '_office.pdf', '_print.xlsx'))
        except Exception as e:
            print(f'xlsx 页面设置注入失败(按原样转换): {e}')
            cur_in = abs_in

    wpscli = _find_wpscli()
    if wpscli:
        sub = 'word2pdf' if ext == '.docx' else 'excel2pdf'
        try:
            if os.path.exists(abs_out):
                os.remove(abs_out)  # 存在时 wpscli 会自动改名, 先删掉
            subprocess.run([wpscli, sub, cur_in, '-o', abs_out],
                           timeout=300, capture_output=True)
            if os.path.exists(abs_out):
                print(f'Office 转换: wpscli {sub}')
                return abs_out
        except Exception as e:
            print(f'wpscli 转换失败: {e}')

    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print('未安装 pywin32, 跳过 Office 转换 (pip install pywin32)')
        return None
    progids = (['Kwps.Application', 'wps.Application', 'Word.Application']
               if ext == '.docx'
               else ['Ket.Application', 'et.Application', 'Excel.Application'])
    pythoncom.CoInitialize()
    app = None
    try:
        for pid in progids:
            try:
                app = win32com.client.DispatchEx(pid)
                print(f'Office 转换: 使用 {pid}')
                break
            except Exception:
                app = None
        if app is None:
            print('未找到可用的 WPS/Office COM 接口')
            return None
        try:
            app.Visible = False
        except Exception:
            pass
        err = None
        for _ in range(2):  # WPS 冷启动/残留实例时第一次调用可能失败, 重试一次
            try:
                if ext == '.docx':
                    doc = app.Documents.Open(abs_in, ReadOnly=True)
                    try:
                        doc.ExportAsFixedFormat(abs_out, 17)
                    except Exception:
                        doc.SaveAs(abs_out, FileFormat=17)  # wdFormatPDF
                    doc.Close(False)
                else:
                    # 注意: ET 的 Workbooks.Open 不支持 ReadOnly 关键字
                    wb = app.Workbooks.Open(cur_in)
                    wb.ExportAsFixedFormat(0, abs_out)    # xlTypePDF
                    wb.Close(False)
                break
            except Exception as e:
                err = e
        if not os.path.exists(abs_out):
            print(f'Office 转换失败: {err}')
            return None
        return abs_out
    except Exception as e:
        print(f'Office 转换失败: {e}')
        return None
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


# ---------------------------------------------------------------------------
# Word (.docx, python-docx) -> 标准 JSON
# ---------------------------------------------------------------------------

def _docx_effective_style(run, para):
    """样式继承链: run 直接格式 -> 段落样式(含 base_style 链) -> 文档默认。

    返回 (字号pt, 加粗, 颜色#rrggbb 或 None, 字体名)。
    """
    def walk(getter):
        v = getter(run.font) if run is not None else None
        if v is not None:
            return v
        st = para.style
        while st is not None:
            v = getter(st.font)
            if v is not None:
                return v
            st = st.base_style
        return None

    size = walk(lambda f: f.size)
    bold = walk(lambda f: f.bold)
    name = walk(lambda f: f.name)
    try:
        color = walk(lambda f: f.color.rgb if f.color and f.color.rgb
                     else None)
    except Exception:
        color = None
    return (size.pt if size is not None else 10.5, bool(bold),
            f'#{color}' if color else None, name)


def _docx_para_role(para):
    """标题层级: 样式名/大纲级别直接判定 (标题1/Heading 1 -> title 等)。"""
    name = (para.style.name or '') if para.style else ''
    low = name.lower()
    if low in ('title', '标题'):
        return 'title'
    for lv, role in ((1, 'title'), (2, 'heading'), (3, 'heading')):
        if f'heading {lv}' in low or f'标题 {lv}' in name \
                or f'标题{lv}' in name:
            return role
    return 'body'


def docx_to_data(path, dpi=96):
    """Word 文档 -> 标准 JSON。

    Word 是流式格式没有固定坐标, 这里做简化排版: 按节设置取页面/页边距,
    段落按行高顺序下排, 表格按列宽/行高换算成格子。dpi 取 96 (屏幕像素),
    JSON 坐标即像素。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    sec = doc.sections[0]
    emu = 12700  # EMU per pt
    page_w = (sec.page_width or 7560310) / emu    # 默认 A4
    page_h = (sec.page_height or 10692000) / emu
    ml = (sec.left_margin or 914400) / emu
    mr = (sec.right_margin or 914400) / emu
    mt = (sec.top_margin or 914400) / emu
    mb = (sec.bottom_margin or 914400) / emu
    s = dpi / 72.0
    content_w = page_w - ml - mr

    texts, cells, rules = [], [], []
    y = mt

    def add_paragraph(p):
        nonlocal y
        text = p.text.strip()
        pf = p.paragraph_format
        runs = [r for r in p.runs if r.text.strip()]
        dom = max(runs, key=lambda r: len(r.text)) if runs else None
        size, bold, color, fname = _docx_effective_style(dom, p)
        spacing = pf.line_spacing
        if spacing is None:
            line_h = size * 1.35
        elif isinstance(spacing, float):
            line_h = size * spacing
        else:
            line_h = spacing / emu
        before = (pf.space_before or 0) and pf.space_before / emu
        after = (pf.space_after or 0) and pf.space_after / emu
        y += before
        if text:
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            align = {WD_ALIGN_PARAGRAPH.CENTER: 'center',
                     WD_ALIGN_PARAGRAPH.RIGHT: 'right'}.get(p.alignment)
            indent = pf.left_indent / emu if pf.left_indent else 0
            texts.append({
                'x': (ml + indent) * s, 'y': (y + size * 0.85) * s,
                'size': size, 'font': fname, 'color': color,
                'bold': bold, 'role': _docx_para_role(p),
                'text': text, 'align': align, 'w': content_w * s,
            })
        y += line_h + after

    def add_table(tbl):
        nonlocal y
        # 列宽: tblGrid 的 dxa (1/20pt); 缺失则均分版心
        grid = tbl._tbl.tblGrid.findall(qn('w:gridCol'))
        widths = [(int(gc.get(qn('w:w'))) / 20 if gc.get(qn('w:w'))
                   else content_w / max(1, len(grid))) for gc in grid]
        ncol = max(1, len(widths))
        xs = [ml]
        for w in widths:
            xs.append(xs[-1] + w)

        # 行高: trHeight (dxa), 缺失按内容行数估算
        rows_y = []
        yy = y
        for row in tbl.rows:
            h = None
            trPr = row._tr.find(qn('w:trPr'))
            if trPr is not None:
                trH = trPr.find(qn('w:trHeight'))
                if trH is not None and trH.get(qn('w:val')):
                    h = int(trH.get(qn('w:val'))) / 20
            if h is None:
                n = max((c.text.count('\n') + 1 for c in row.cells),
                        default=1)
                h = max(16.0, n * 10.5 * 1.35)
            rows_y.append((yy, h))
            yy += h

        # 表格边框: tblBorders (sz 单位 1/8pt); 无定义时 Table Grid 默认全框
        borders = {}
        tblPr = tbl._tbl.tblPr
        tb = tblPr.find(qn('w:tblBorders')) if tblPr is not None else None
        if tb is not None:
            for side in ('top', 'left', 'bottom', 'right',
                         'insideH', 'insideV'):
                el = tb.find(qn(f'w:{side}'))
                if el is None or el.get(qn('w:val')) in ('none', 'nil'):
                    borders[side] = None
                else:
                    borders[side] = int(el.get(qn('w:sz'), '4')) / 8
        else:
            grid_style = tbl.style and 'grid' in (tbl.style.name or '').lower()
            borders = {k: (0.5 if grid_style else None)
                       for k in ('top', 'left', 'bottom', 'right',
                                 'insideH', 'insideV')}

        def rule(x0, y0, x1, y1, wpt):
            if wpt:
                rules.append({'x0': x0 * s, 'y0': y0 * s,
                              'x1': x1 * s, 'y1': y1 * s,
                              'width': wpt * s})

        # 逐行逐格: 直接读 tr.tc_lst (row.cells 会把 vMerge 纵向合并
        # 解析成根格内容, 掩盖延续格); gridSpan 横并, vMerge 纵并
        from docx.table import _Cell
        tbl_start = len(cells)
        vopen = {}  # col -> cells 列表里跨行来源格的索引
        for r, row in enumerate(tbl.rows):
            c = 0
            for tc in row._tr.tc_lst:
                span = tc.grid_span or 1
                vm = tc.vMerge
                if c >= ncol:
                    break
                if vm == 'continue':
                    if c in vopen:
                        cells[vopen[c]]['rowspan'] += 1
                    c += span
                    continue
                cell = _Cell(tc, tbl)
                text = cell.text.strip()
                ent = {'row': r, 'col': c, 'rowspan': 1, 'colspan': span,
                       'x': xs[c] * s, 'y': rows_y[r][0] * s,
                       'w': (xs[min(c + span, ncol)] - xs[c]) * s,
                       'h': 0,  # 行高可能因纵并后延, 最后统一算
                       '_r0': r,
                       'borders': None, 'text': text}
                para0 = next((p for p in cell.paragraphs if p.text.strip()),
                             None)
                if para0 is not None:
                    runs = [rn for rn in para0.runs if rn.text.strip()]
                    dom = (max(runs, key=lambda rn: len(rn.text))
                           if runs else None)
                    size, bold, color, fname = _docx_effective_style(
                        dom, para0)
                    ent['size'] = size
                    ent['font'] = fname
                    if bold:
                        ent['bold'] = True
                vopen.pop(c, None)
                if vm == 'restart':
                    vopen[c] = len(cells)
                cells.append(ent)
                c += span

        # 纵并结束后算每格真实高度
        for ent in cells:
            if ent.get('_r0') is not None and ent['h'] == 0:
                r0 = ent['_r0']
                ent['h'] = sum(rows_y[i][1]
                               for i in range(r0, min(r0 + ent['rowspan'],
                                                      len(rows_y)))) * s
                del ent['_r0']

        # 边框 -> 矢量线段 (合并区域内部的分隔线不画)
        tbl_cells = cells[tbl_start:]
        hskip = set()  # (行边界 r, 列 c): 纵向合并跨过
        vskip = set()  # (行 r, 列边界 c): 横向合并跨过
        for ent in tbl_cells:
            for rr in range(ent['row'] + 1, ent['row'] + ent['rowspan']):
                for cc in range(ent['col'], ent['col'] + ent['colspan']):
                    hskip.add((rr, cc))
            for rr in range(ent['row'], ent['row'] + ent['rowspan']):
                for cc in range(ent['col'] + 1, ent['col'] + ent['colspan']):
                    vskip.add((rr, cc))
        nrow = len(rows_y)
        for r in range(nrow + 1):
            yb = rows_y[r][0] if r < nrow else yy
            wpt = borders['top'] if r == 0 else (
                borders['bottom'] if r == nrow else borders['insideH'])
            if not wpt:
                continue
            c = 0
            while c < ncol:
                if (r, c) in hskip:
                    c += 1
                    continue
                c0 = c
                while c + 1 < ncol and (r, c + 1) not in hskip:
                    c += 1
                rule(xs[c0], yb, xs[c + 1], yb, wpt)
                c += 1
        for c in range(ncol + 1):
            wpt = borders['left'] if c == 0 else (
                borders['right'] if c == ncol else borders['insideV'])
            if not wpt:
                continue
            r = 0
            while r < nrow:
                if (r, c) in vskip:
                    r += 1
                    continue
                r0 = r
                while r + 1 < nrow and (r + 1, c) not in vskip:
                    r += 1
                rule(xs[c], rows_y[r0][0], xs[c],
                     rows_y[r][0] + rows_y[r][1], wpt)
                r += 1
        y = yy + 12  # 表格后留一点间距

    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            add_paragraph(Paragraph(child, doc))
        elif child.tag == qn('w:tbl'):
            add_table(Table(child, doc))

    # 清理辅助字段
    for ent in cells:
        ent.pop('_r0', None)
        if ent['borders'] is None:
            del ent['borders']

    height = max(page_h, y + mb)
    return {'dpi': dpi, 'width': int(page_w * s), 'height': int(height * s),
            'cells': cells, 'texts': texts, 'rules': rules}


# ---------------------------------------------------------------------------
# Excel (.xlsx, openpyxl) -> 标准 JSON
# ---------------------------------------------------------------------------

def xlsx_to_data(path, dpi=96):
    """Excel 工作簿 -> 标准 JSON。

    逐 sheet 遍历: 值/公式/样式/合并区域分开记录; 多个 sheet 纵向拼接,
    之间留 40px 间距并标注 sheet 名。列宽(字符单位)和行高(pt)换算成像素。
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb_f = openpyxl.load_workbook(path)                  # 公式
    wb_v = openpyxl.load_workbook(path, data_only=True)  # 缓存值
    s = dpi / 72.0
    bwidth = {'hair': 0.5, 'thin': 1.0, 'medium': 2.0, 'thick': 3.0,
              'double': 3.0, 'dashed': 1.0, 'dotted': 0.5}

    cells, texts, rules = [], [], []
    by_addr = {}
    width = 0
    y_off = 0
    row_base = 0  # 多 sheet 纵向拼接时行号全局唯一化, 供角色/字段按网格邻接
    for name in wb_f.sheetnames:
        ws = wb_f[name]
        wsv = wb_v[name]
        if ws.max_row is None:
            continue
        if len(wb_f.sheetnames) > 1:
            texts.append({'x': 0, 'y': y_off + 18 * s, 'size': 14,
                          'font': None, 'color': None, 'bold': True,
                          'role': 'heading', 'text': f'[{name}]',
                          'align': None, 'w': 0})
            y_off += int(30 * s)

        col_x = [0]
        for i in range(1, ws.max_column + 1):
            cd = ws.column_dimensions.get(get_column_letter(i))
            w = cd.width if cd and cd.width else 8.43
            col_x.append(col_x[-1] + int(round(w * 7 + 5)))  # 96dpi 近似
        row_y = [0]
        for i in range(1, ws.max_row + 1):
            rd = ws.row_dimensions.get(i)
            h = rd.height if rd and rd.height else 13.2
            row_y.append(row_y[-1] + int(round(h * s)))

        merged = {}   # 被合并覆盖的位置 -> 锚点 (r, c)
        spans = {}    # 锚点 -> (rowspan, colspan)
        for rg in ws.merged_cells.ranges:
            spans[(rg.min_row, rg.min_col)] = (
                rg.max_row - rg.min_row + 1, rg.max_col - rg.min_col + 1)
            for rr in range(rg.min_row, rg.max_row + 1):
                for cc in range(rg.min_col, rg.max_col + 1):
                    if (rr, cc) != (rg.min_row, rg.min_col):
                        merged[(rr, cc)] = (rg.min_row, rg.min_col)

        seen_rules = {}

        def rule(x0, y0, x1, y1, wpx):
            key = (round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1))
            if key not in seen_rules or seen_rules[key] < wpx:
                seen_rules[key] = wpx

        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                if (r, c) in merged:
                    continue
                cell = ws.cell(row=r, column=c)
                vcell = wsv.cell(row=r, column=c)
                value = cell.value
                formula = None
                if isinstance(value, str) and value.startswith('='):
                    formula = value
                    value = vcell.value  # 显示用缓存计算值
                rowspan, colspan = spans.get((r, c), (1, 1))
                x0 = col_x[c - 1]
                y0 = row_y[r - 1] + y_off
                x1 = col_x[min(c - 1 + colspan, len(col_x) - 1)]
                y1 = row_y[min(r - 1 + rowspan, len(row_y) - 1)] + y_off
                ent = {'row': r - 1 + row_base, 'col': c - 1,
                       'rowspan': rowspan, 'colspan': colspan,
                       'x': x0, 'y': y0, 'w': x1 - x0, 'h': y1 - y0,
                       'borders': None,
                       'addr': f'{name}!{get_column_letter(c)}{r}',
                       'text': '' if value is None else str(value)}
                if formula:
                    ent['formula'] = formula
                vv = vcell.value
                if vv is not None:
                    ent['value'] = _jsonable(vv)
                nf = cell.number_format
                if nf and nf != 'General':
                    ent['number_format'] = nf
                if value is not None and str(value).strip():
                    room = sum(1 for k in range(
                        c + 1, min(c + colspan, ws.max_column) + 1)
                        if wsv.cell(row=r, column=k).value is None)
                    if room:
                        ent['value_room'] = room
                f = cell.font
                if f:
                    if f.sz:
                        ent['size'] = f.sz
                    if f.b:
                        ent['bold'] = True
                    if f.name:
                        ent['font'] = f.name
                    try:
                        if f.color and f.color.rgb \
                                and isinstance(f.color.rgb, str):
                            ent['color'] = '#' + f.color.rgb[-6:]
                    except Exception:
                        pass
                al = cell.alignment
                if al and al.horizontal in ('left', 'right', 'center'):
                    ent['align'] = al.horizontal
                cells.append(ent)
                by_addr[ent['addr']] = ent

        # 边框 -> 矢量线段: 单独遍历所有格(含被合并覆盖的格, 其边上
        # 可能定义了框线), 共享边去重取最粗; 同时记录每格四边有无,
        # 供角色标注 (表体可填判定) 使用
        cell_borders = {}
        for row in ws.iter_rows():
            for cell in row:
                bd = cell.border
                if not bd:
                    continue
                r, c = cell.row, cell.column
                x0 = col_x[c - 1]
                y0 = row_y[r - 1] + y_off
                x1 = col_x[min(c, len(col_x) - 1)]
                y1 = row_y[min(r, len(row_y) - 1)] + y_off
                sides = {}
                for side, seg in (
                        ('top', (x0, y0, x1, y0)),
                        ('bottom', (x0, y1, x1, y1)),
                        ('left', (x0, y0, x0, y1)),
                        ('right', (x1, y0, x1, y1))):
                    b = getattr(bd, side, None)
                    has = bool(b and b.style in bwidth)
                    sides[side] = has
                    if has:
                        rule(*seg, bwidth[b.style])
                cell_borders[(r, c)] = sides
        for (r, c), sides in cell_borders.items():
            ent = by_addr.get(f'{name}!{get_column_letter(c)}{r}')
            if ent is not None:
                ent['borders'] = sides

        rules.extend({'x0': k[0], 'y0': k[1], 'x1': k[2], 'y1': k[3],
                      'width': w} for k, w in seen_rules.items())
        width = max(width, col_x[-1])
        y_off = row_y[-1] + y_off + int(40 * s)
        row_base += ws.max_row

    for ent in cells:
        if ent.get('borders') is None:
            ent.pop('borders', None)
    return {'dpi': dpi, 'width': width,
            'height': y_off - int(40 * s) if y_off else 0,
            'cells': cells, 'texts': texts, 'rules': rules}


# ---------------------------------------------------------------------------
# 可回填表单模板: 语义对齐 (addr) + 角色标注 + 字段配对
#
# 回填契约: 按 fields/table_fields 按键名匹配找到目标 addr, 把对应 cell 的
# text 改为新值, 并删除该 cell 的 lines 字段 (lines 是 PDF 文字层的逐行缓存,
# 不删的话旧值会被渲染出来)。text 是唯一回填入口, 几何与边框不许动。
# ---------------------------------------------------------------------------

def _norm_text(s):
    """去掉所有空白后的文本, 用于跨源匹配 (PDF 文字层 vs openpyxl 值)。"""
    return ''.join(str(s).split())


def _jsonable(v):
    """把 openpyxl 单元格值转成 JSON 可序列化形态。"""
    import datetime
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    return str(v)


def _group_free_text_lines(texts):
    """把游离文字按基线 y 聚成视觉行。

    返回 [{'y', 'items': [{'idx', 'x', 'y', 'text'}]}], items 按 x 升序,
    idx 为在 texts 列表中的下标 (供 #T{n} 寻址)。
    """
    items = sorted(({'idx': i, 'x': t['x'], 'y': t['y'], 'text': t['text']}
                    for i, t in enumerate(texts)),
                   key=lambda it: (it['y'], it['x']))
    lines = []
    for it in items:
        if lines and abs(it['y'] - lines[-1]['y']) <= 6:
            lines[-1]['items'].append(it)
        else:
            lines.append({'y': it['y'], 'items': [it]})
    return lines


def _table_rows_set(cells, header_row):
    """表体区行集合: 表头行之下、行标签列 (非空格最多的列) 之左无文本的连续行。

    表体区之后的行 (检查结论/审核人等表单尾区) 返回时不包含,
    供角色/字段区分"表格"与"表单尾部"。
    """
    by_row = {}
    for c in cells:
        by_row.setdefault(c['row'], []).append(c)
    counts = {}
    for c in cells:
        if c['row'] > header_row and (c.get('text') or '').strip():
            counts[c['col']] = counts.get(c['col'], 0) + 1
    table_rows = set()
    if counts:
        label_col = max(counts, key=counts.get)
        r = header_row + 1
        while r in by_row:
            row_cells = by_row[r]
            if any((c.get('text') or '').strip()
                   for c in row_cells if c['col'] < label_col):
                break
            # 表单尾区标签行 (检查结论/审核人): 行首短文本无符号前缀,
            # 且右侧整行无内容 -> 不是表体行, 表体到此为止
            def _tail_label(c):
                t = (c.get('text') or '').replace('\n', '').strip()
                return bool(t) and len(t) <= 12 \
                    and t[0] not in _NO_LABEL_PREFIX \
                    and not any((x.get('text') or '').strip()
                                for x in row_cells if x['col'] > c['col'])
            if any(_tail_label(c) for c in row_cells
                   if c['col'] < max(label_col, 1)):
                break
            table_rows.add(r)
            r += 1
    return table_rows


def augment_xlsx_semantics(page_data, wb_f, wb_v, max_candidates=50):
    """把一页探测结果对齐到原生 workbook, 给格子/游离文字赋稳定地址 addr。

    对齐用纯文本锚点投票: 探测格文本(去空白)在原生各 sheet 中反查,
    每个同名位置产生候选偏移 (dr,dc), 得票 Top 做全量校验取匹配最多者;
    阈值 max(3, 0.6×非空探测格), 跨 sheet 取总分最高, 不达标该页跳过
    (JSON 退回纯视觉, 不产生错误地址)。命中后逐格赋
    addr ("Sheet1!B2") 及 value/formula/number_format,
    字体样式 (font_size/font_bold/font_name/font_color/font_align),
    游离文字行能整行匹配原生格的也补 addr。
    字号字体对回填渲染是必需信息: 空值格 PDF 里没有文字层, 只能从
    原生字体取 (PDF 文字层的实际打印字号优先, 见 build_index.ref)。
    """
    from openpyxl.utils import get_column_letter

    sheets = []
    for name in wb_f.sheetnames:
        ws, wsv = wb_f[name], wb_v[name]
        disp, sem, fonts, aligns = {}, {}, {}, {}
        for row in ws.iter_rows():
            for cell in row:
                fnt = cell.font
                if fnt is not None:
                    fe = {}
                    if fnt.sz:
                        fe['size'] = float(fnt.sz)
                    if fnt.b:
                        fe['bold'] = True
                    if fnt.name:
                        fe['name'] = fnt.name
                    try:
                        if fnt.color and fnt.color.rgb \
                                and isinstance(fnt.color.rgb, str):
                            fe['color'] = '#' + fnt.color.rgb[-6:]
                    except Exception:
                        pass
                    if fe:
                        fonts[(cell.row, cell.column)] = fe
                al = cell.alignment
                if al is not None and al.horizontal in (
                        'left', 'right', 'center'):
                    aligns[(cell.row, cell.column)] = al.horizontal
                v = cell.value
                nf = cell.number_format
                if v is None and (not nf or nf == 'General'):
                    continue  # 无值无格式的真空格, 不需要语义记录
                r, c = cell.row, cell.column
                ent = {}
                if v is not None:
                    vv = wsv.cell(row=r, column=c).value
                    if vv is not None:
                        # 无缓存的公式 (如 openpyxl 直接生成未经过 Excel
                        # 重存) 不进 disp: PDF 里是计算结果, 匹配不上
                        disp[(r, c)] = _norm_text(vv)
                        ent['value'] = _jsonable(vv)
                    if isinstance(v, str) and v.startswith('='):
                        ent['formula'] = v
                if nf and nf != 'General':
                    ent['number_format'] = nf
                sem[(r, c)] = ent
        sheets.append((name, disp, sem, ws.max_row, ws.max_column, fonts,
                       aligns))

    anchors = [(_norm_text(c.get('text') or ''), c['row'], c['col'])
               for c in page_data['cells']]
    anchors = [a for a in anchors if len(a[0]) >= 2]
    if not anchors:
        return
    votes = {}
    for t, row, col in anchors:
        for name, disp, _, _, _, _, _ in sheets:
            for (r, c), dt in disp.items():
                if dt == t:
                    key = (name, r - 1 - row, c - 1 - col)
                    votes[key] = votes.get(key, 0) + 1
    if not votes:
        print('语义对齐: 无锚点命中, 该页跳过')
        return

    best = None  # (matches, total, name, dr, dc)
    for (name, dr, dc), _ in sorted(votes.items(), key=lambda kv: -kv[1]
                                    )[:max_candidates]:
        disp = next(d for n, d, _, _, _, _, _ in sheets if n == name)
        matches = total = 0
        for c in page_data['cells']:
            t = _norm_text(c.get('text') or '')
            if len(t) < 2:
                continue
            total += 1
            if disp.get((c['row'] + dr + 1, c['col'] + dc + 1)) == t:
                matches += 1
        if best is None or matches > best[0]:
            best = (matches, total, name, dr, dc)
    matches, total, name, dr, dc = best
    need = max(3, -(-total * 3 // 5))
    if matches < need:
        print(f'语义对齐: 最佳匹配 {matches}/{total} 低于阈值 {need}, 该页跳过')
        return
    print(f'语义对齐: {name} 偏移=({dr},{dc}) 匹配 {matches}/{total}')

    _, _, sem, max_row, max_col, fonts, aligns = next(
        s for s in sheets if s[0] == name)
    for c in page_data['cells']:
        r, cc = c['row'] + dr + 1, c['col'] + dc + 1
        if not (1 <= r <= max_row and 1 <= cc <= max_col):
            continue
        # 使用范围内每格都给 addr (空格是主要回填目标, 不能因为原生为空而缺地址)
        c['addr'] = f"{name}!{get_column_letter(cc)}{r}"
        c.update(sem.get((r, cc), {}))
        fnt = fonts.get((r, cc))
        if fnt:
            if fnt.get('size') is not None:
                c.setdefault('font_size', fnt['size'])
            if fnt.get('bold'):
                c.setdefault('font_bold', True)
            if fnt.get('name'):
                c.setdefault('font_name', fnt['name'])
            if fnt.get('color'):
                c.setdefault('font_color', fnt['color'])
        al = aligns.get((r, cc))
        if al:
            c.setdefault('font_align', al)
        # 格内右侧留白: colspan 范围内除锚点外还有原生空格 -> 值可追加
        if (c.get('text') or '').strip():
            room = sum(1 for k in range(cc + 1, min(cc + c.get('colspan', 1),
                                                   max_col) + 1)
                       if (r, k) not in disp)
            if room:
                c['value_room'] = room

    texts = page_data.get('texts') or []
    for ln in _group_free_text_lines(texts):
        items = ln['items']
        joined = ''.join(it['text'] for it in items)
        # 逐字符建立 (归一化位置 -> 原文位置/所属 item) 映射
        cmap, npos, nmap = [], '', []
        for k, it in enumerate(items):
            cmap.extend([k] * len(it['text']))
        for ci, ch in enumerate(joined):
            n = _norm_text(ch)
            if n:
                npos += n
                nmap.append(ci)
        if len(npos) < 2:
            continue
        # 整行优先; 整行配不上时按"最长前缀逐段消费"匹配
        # (一行跨多个原生格: "监理单位：  编 号：JB12-T-1")
        start = 0
        while start < len(npos):
            hit = None
            for (r, c), dt in disp.items():
                if dt and npos.startswith(dt, start) and \
                        (hit is None or len(dt) > len(hit[2])):
                    hit = (r, c, dt)
            if hit is None:
                break
            r, c, dt = hit
            lo, hi = nmap[start], nmap[start + len(dt) - 1]
            a = f"{name}!{get_column_letter(c)}{r}"
            for k in range(cmap[lo], cmap[hi] + 1):
                texts[items[k]['idx']]['addr'] = a
            start += len(dt)


_NO_LABEL_PREFIX = tuple('0123456789△▲□■＋±+-≤≥<>＝=%')


def assign_roles(data):
    """给每个格子标注角色 title|header|label|value|data 并标 fillable。

    纯启发式 (按序判定, 先命中先生效), 供回填定位; 结果打印统计,
    配合 overlay.png 人工核对。表头行之下的空边框格一律可填 (含备用空行)。
    """
    cells = data['cells']
    if not cells:
        return
    by_grid = {(c['row'], c['col']): c for c in cells}
    page_cx = data['width'] / 2

    def text_of(c):
        return (c.get('text') or '').replace('\n', '').strip()

    def cell_sizes(c):
        return [l['size'] for l in c.get('lines', []) if 'size' in l]

    sizes = sorted(s for c in cells for s in cell_sizes(c))
    body = sizes[len(sizes) // 2] if sizes else 10.0

    # 1) title: 格内大字且接近页水平居中 (四边全框的格是表格字段, 排除)
    for c in cells:
        b = c.get('borders') or {}
        if all(b.get(s) for s in ('top', 'bottom', 'left', 'right')):
            continue
        szs = cell_sizes(c)
        if szs and max(szs) >= body * 1.3 \
                and abs(c['x'] + c['w'] / 2 - page_cx) < data['width'] * 0.12:
            c['role'] = 'title'

    # 2) header: 某行非空格 >=3、全部短文本, 且 >=2 个表头列的下一行同列
    #    有内容 (真表头直接压在数据列上; "抽检人/抽检时间/年月日"这类
    #    信息行下面是空行或另一表头, 不会误标)
    header_rows = set()
    by_row = {}
    for c in cells:
        by_row.setdefault(c['row'], []).append(c)
    for r, row_cells in sorted(by_row.items()):
        non_empty = [text_of(c) for c in row_cells if text_of(c)]
        if len(non_empty) < 3 or not all(len(t) <= 10 for t in non_empty):
            continue
        hcols = {c['col'] for c in row_cells if text_of(c)}
        filled_below = sum(
            1 for hc in hcols
            if (nxt := by_grid.get((r + 1, hc))) is not None
            and text_of(nxt))
        if filled_below < 2:
            continue
        for c in row_cells:
            if text_of(c):
                c['role'] = 'header'
        header_rows.add(r)

    header_row = min(header_rows) if header_rows else None

    # 2.5) 表体区 (检查结论/审核人等表单尾区不在其中)
    table_rows = _table_rows_set(cells, header_row) \
        if header_row is not None else set()

    # 3) 表体区的空边框格: 可填 (含专用填写行/备用空行)
    for c in cells:
        if c['row'] in table_rows and not text_of(c) \
                and any((c.get('borders') or {}).values()):
            c['role'] = 'value'
            c['fillable'] = True

    # 4) label: 非空短文本, 不以数值/符号开头, 不在表体区,
    #    右邻为空(或下划线占位)或下邻为空
    for c in cells:
        if c.get('role') or c['row'] in table_rows:
            continue
        t = text_of(c)
        if not t or len(t) > 12 or t[0] in _NO_LABEL_PREFIX:
            continue
        right = by_grid.get((c['row'], c['col'] + c['colspan']))
        below = by_grid.get((c['row'] + c['rowspan'], c['col']))
        right_empty = right is not None and (
            not text_of(right) or '___' in (right.get('text') or ''))
        below_empty = below is not None and not text_of(below)
        if right_empty or below_empty:
            c['role'] = 'label'

    # 5) value: 空格, 左邻或上邻是 label (表体区已在 3 覆盖)
    for c in cells:
        if c.get('role') or text_of(c):
            continue
        left = by_grid.get((c['row'], c['col'] - 1))
        above = by_grid.get((c['row'] - 1, c['col']))
        if (left is not None and left.get('role') == 'label') or \
           (above is not None and above.get('role') == 'label'):
            c['role'] = 'value'
            c['fillable'] = True

    # 6) data: 其余非空格
    for c in cells:
        if not c.get('role') and text_of(c):
            c['role'] = 'data'

    roles = {}
    for c in cells:
        r = c.get('role') or '-'
        roles[r] = roles.get(r, 0) + 1
    print(f'角色标注: {dict(sorted(roles.items()))}')


_COLON_RE = None  # 延迟编译, 见 extract_fields


def extract_fields(data):
    """从角色标注中产出可回填字段: fields (标签->值/内联) + table_fields。

    - 格内联冒号: 整格文本形如 "编 号：xxx" -> inline 字段, prefix 为冒号前;
    - 标签->值对: label 角色格优先配右邻 (下划线占位格也算), 下邻次之;
    - 游离文字行: 按基线聚行后同样匹配冒号模式, 用 addr 或 #T{n} 寻址,
      segs 为该行涉及的 texts 下标, value_seg 为值起始的下标位置;
    - 表体行: 表头行之下, 行标签列 = 非空格最多的列; 每个行标签下,
      只收集"标签行该列为空"的表头列, 优先取下一行 (专用填写行) 的地址。
    """
    import re
    global _COLON_RE
    if _COLON_RE is None:
        _COLON_RE = re.compile(r'^(.*?[:：])(.*)$', re.S)

    cells = data['cells']
    by_grid = {(c['row'], c['col']): c for c in cells}

    def text_of(c):
        return (c.get('text') or '').replace('\n', '').strip()

    fields, table_fields = [], []

    # 1) 格内联冒号
    for c in cells:
        t = (c.get('text') or '').strip()
        m = _COLON_RE.match(t)
        if m and len(t) <= 20 and c.get('addr'):
            name = re.sub(r'\s+', '', re.sub(r'\s*[:：]\s*$', '',
                                             m.group(1)))
            if name:
                fields.append({'name': name, 'label': c['addr'],
                               'inline': True, 'prefix': m.group(1)})

    # 2) 标签->值对
    for c in cells:
        if c.get('role') != 'label' or not c.get('addr'):
            continue
        name = re.sub(r'\s+', '', re.sub(r'\s*[:：]\s*$', '', text_of(c)))
        if not name:
            continue
        target = by_grid.get((c['row'], c['col'] + c['colspan']))
        if target is not None and text_of(target) \
                and '___' not in (target.get('text') or ''):
            target = None
        if target is None:
            below = by_grid.get((c['row'] + c['rowspan'], c['col']))
            if below is not None and not text_of(below):
                target = below
        if target is None or not target.get('addr'):
            continue
        target['fillable'] = True
        target.setdefault('role', 'value')
        entry = {'name': name, 'label': c['addr'], 'value': target['addr']}
        if '___' in (target.get('text') or ''):
            entry['underscore'] = True  # 回填时把下划线段替换为值
        fields.append(entry)

    # 2.5) 格内右侧留白型: 标签与值探测成同一格, 但原生 colspan 内锚点
    #      右侧还有空格 (value_room, 对齐阶段记录) -> 值追加写在同一格
    #      标签之后 (如 "抽检时间" 后的大片空格)
    header_row0 = min((c['row'] for c in cells
                       if c.get('role') == 'header'), default=None)
    table_rows = _table_rows_set(cells, header_row0) \
        if header_row0 is not None else set()
    claimed = {a for f in fields for a in (f.get('label'), f.get('value'))}
    for c in cells:
        if c.get('role') in ('header', 'title', 'value') \
                or c['row'] in table_rows \
                or not c.get('addr') or c['addr'] in claimed:
            continue
        t = text_of(c)
        if not t or len(t) > 12 or t[0] in _NO_LABEL_PREFIX:
            continue
        if not c.get('value_room'):
            continue
        c['fillable'] = True
        # 左邻是标签样文本 -> 本格是它的值格, 已有文字视为占位, 整格替换
        # (占位文字要求短, 避免把 "基本要求" 配到后面的长 checklist 格)
        left = by_grid.get((c['row'], c['col'] - 1))
        lt = text_of(left) if left is not None else ''
        if lt and len(t) <= 8 and len(lt) <= 12 \
                and lt[0] not in _NO_LABEL_PREFIX \
                and left.get('role') in (None, 'data') \
                and left.get('addr') and left['addr'] not in claimed:
            fields.append({'name': re.sub(r'\s+', '', lt),
                           'label': left['addr'], 'value': c['addr'],
                           'replace': True})
            claimed.add(left['addr'])
            continue
        fields.append({'name': re.sub(r'\s+', '', t),
                       'label': c['addr'], 'value': c['addr'],
                       'inline': True, 'prefix': t, 'append': True})

    # 3) 游离文字行: 一行可含多个"标签：值"字段 (按冒号扫描, 如
    #    "监理单位：    编  号：JB12-4.2.2-")
    texts = data.get('texts') or []
    label_chars = re.compile(r'([\w\u4e00-\u9fff（）()·、\s]{1,14}?)\s*[:：]')
    for ln in _group_free_text_lines(texts):
        items = ln['items']
        joined = ''.join(it['text'] for it in items)
        if len(_norm_text(joined)) > 40:  # 长段落不是字段行
            continue
        matches = list(label_chars.finditer(joined))
        if not matches:
            continue

        def seg_at(pos):
            acc = 0
            for k, it in enumerate(items):
                acc += len(it['text'])
                if pos < acc:
                    return k
            return len(items) - 1

        for k, m in enumerate(matches):
            name = re.sub(r'\s+', '', m.group(1))
            if not name or len(name) > 12 or name[0] in _NO_LABEL_PREFIX \
                    or '□' in name:
                continue
            value_start = m.end()
            value_end = (matches[k + 1].start() if k + 1 < len(matches)
                         else len(joined))
            ls = seg_at(m.start())
            entry = {'name': name,
                     'label': texts[items[ls]['idx']].get('addr')
                              or f"#T{items[ls]['idx']}",
                     'inline': True, 'prefix': m.group(0)}
            value_text = joined[value_start:value_end].strip()
            if value_text:
                vs = seg_at(value_start)
                entry['value'] = (texts[items[vs]['idx']].get('addr')
                                  or f"#T{items[vs]['idx']}")
                entry['value_seg'] = vs
            fields.append(entry)

    # 4) 表体行字段
    header_row = min((c['row'] for c in cells if c.get('role') == 'header'),
                     default=None)
    if header_row is not None:
        headers = {}
        for c in cells:
            if c['row'] == header_row and text_of(c):
                headers.setdefault(c['col'], text_of(c))
        table_rows = _table_rows_set(cells, header_row)
        body = [c for c in cells if c['row'] in table_rows]
        counts = {}
        for c in body:
            if text_of(c):
                counts[c['col']] = counts.get(c['col'], 0) + 1
        if counts:
            label_col = max(counts, key=counts.get)
            for c in sorted(body, key=lambda x: (x['row'], x['col'])):
                t = text_of(c)
                if c['col'] != label_col or not t:
                    continue
                entry = {'row_label': t, 'cells': {}}
                if c.get('addr'):
                    entry['row'] = c['addr']
                for hc, hname in headers.items():
                    if hc == label_col:
                        continue
                    same = by_grid.get((c['row'], hc))
                    if same is None or text_of(same):
                        continue  # 该列此行已有内容 (如规定值), 不是填写列
                    for rr in (c['row'] + 1, c['row']):
                        tc = by_grid.get((rr, hc))
                        if tc is None or text_of(tc) or not tc.get('addr'):
                            continue
                        if rr == c['row'] + 1:
                            # 下一行已是新条目 -> 本行该列的留空不属于本条目;
                            # 本格跨行 (rowspan>1) 时回退到本格 (合格率列
                            # 纵向合并, 值就写在合并格里)
                            lbl = by_grid.get((rr, label_col))
                            if lbl is not None and text_of(lbl):
                                continue
                        entry['cells'][hname] = tc['addr']
                        break
                if entry['cells']:
                    table_fields.append(entry)

    data['fields'] = fields
    data['table_fields'] = table_fields
    print(f'字段配对: fields {len(fields)} 个, '
          f'table_fields {len(table_fields)} 行')


def build_form_view(data):
    """把 fields/table_fields 整理成可索引的 form 视图, 嵌入 data['form']。

    为后续建立索引设计, 不是简单组合:
    - 每个字段是一条扁平文档: name/mode/addr/几何坐标全在同一层,
      索引导入时无需展开嵌套对象;
    - 每条有全局唯一 id (F01..), 重名字段 (如两个"年月日") 按 id 区分,
      name 只做匹配键;
    - 同名/易混字段带 context: 同行最近左侧标签 (如 年月日@抽检时间),
      供检索时消歧;
    - 值格几何 (x/y/w/h, 像素, 基于 data['dpi']) 平铺在字段上,
      索引命中即可直接渲染, 不必回 cells 查表。
    mode 语义:
      cell    值写入空格 (改 value_addr 对应格 text)
      append  值追加在同行标签之后 (label/value 同格)
      replace 整格替换 (占位文字视为示例值)
      inline  同行冒号内联, 替换 prefix 之后的文本
    游离文字的值宽按字号估算 (estimated=true)。渲染层 cells/texts/rules
    保持不变, json_to_pdf / index.html 照常消费。
    """
    by_addr = {}
    cells = data['cells']
    for c in cells:
        if c.get('addr'):
            by_addr[c['addr']] = c
    for i, t in enumerate(data.get('texts') or []):
        by_addr[t['addr'] if t.get('addr') else f'#T{i}'] = t
    by_grid = {(c['row'], c['col']): c for c in cells}

    def _flat(addr):
        """addr -> 扁平几何/文本属性 (px); 游离文字宽度按字号估算。"""
        ent = by_addr.get(addr)
        if ent is None:
            return {'addr': addr}
        d = {'addr': addr,
             'x': round(ent['x'], 1), 'y': round(ent['y'], 1),
             'w': round(ent.get('w') or 0, 1),
             'h': round(ent.get('h') or 0, 1)}
        if ent.get('text'):
            d['text'] = ent['text']
        if ent.get('size'):
            d['size'] = round(ent['size'], 2)
        if ent.get('fillable'):
            d['fillable'] = True
        if 'w' not in ent and ent.get('size'):   # 游离文字: 无实体宽
            n = len(''.join(str(ent.get('text') or '').split()))
            d['w'] = round(n * ent['size'] * data.get('dpi', 96) / 72.0, 1)
            d['estimated'] = True
        return d

    def _context(ent):
        """值格同行最近左侧的标签样文本 (同名字段消歧用)。"""
        if ent is None or (ent.get('row') is None):
            return None, None
        best = None
        for c in cells:
            if c['row'] != ent.get('row') or c['col'] >= ent.get('col', 0):
                continue
            t = (c.get('text') or '').replace('\n', '').strip()
            if t and len(t) <= 12 and t[0] not in _NO_LABEL_PREFIX:
                if best is None or c['col'] > best['col']:
                    best = c
        if best is None:
            return None, None
        return (best.get('text') or '').replace('\n', '').strip(), \
            best.get('addr')

    form_fields = []
    for i, f in enumerate(data.get('fields') or []):
        mode = ('inline' if f.get('inline') else
                'replace' if f.get('replace') else
                'append' if f.get('append') else 'cell')
        doc = {'id': f'F{i + 1:02d}', 'name': f['name'], 'mode': mode}
        label_ent = by_addr.get(f.get('label') or '')
        value_ent = by_addr.get(f.get('value') or '')
        if label_ent is not None:
            doc['label_addr'] = f['label']
            if label_ent.get('text'):
                doc['label_text'] = label_ent['text']
        if value_ent is not None and f.get('value'):
            v = _flat(f['value'])
            doc['value_addr'] = v.pop('addr')
            doc.update(v)
        for k in ('prefix', 'underscore', 'value_seg'):
            if k in f:
                doc[k] = f[k]
        ctx_text, ctx_addr = _context(value_ent or label_ent)
        if ctx_text and ctx_text != f['name']:
            doc['context'] = ctx_text
            if ctx_addr:
                doc['context_addr'] = ctx_addr
        form_fields.append(doc)

    form = {'meta': {'dpi': data.get('dpi'), 'width': data.get('width'),
                     'height': data.get('height')},
            'fields': form_fields}
    titles = [t for t in data.get('texts') or [] if t.get('role') == 'title']
    if titles:
        form['meta']['title'] = max(titles, key=lambda t: t.get('size', 0)) \
            ['text'].strip()

    rows = data.get('table_fields') or []
    if rows:
        headers = []
        for r in rows:
            for h in r['cells']:
                if h not in headers:
                    headers.append(h)
        table_rows = []
        for i, r in enumerate(rows):
            row_doc = {'id': f'R{i + 1:02d}', 'label': r['row_label']}
            if r.get('row'):
                row_doc['label_addr'] = r['row']
            row_doc['cells'] = {h: _flat(a) for h, a in r['cells'].items()}
            table_rows.append(row_doc)
        form['table'] = {'headers': headers, 'rows': table_rows}
    data['form'] = form
    print(f'form 视图: 字段 {len(form_fields)} 个'
          + (f', 表体 {len(rows)} 行' if rows else ''))


# ---------------------------------------------------------------------------
# 可填充索引 index: by_label + header_fields + rows (+ cells + report)
#
# 表头行 (role=header, 含"规定值"类列名的行) 为分界线, 构建时自动识别:
#   表头行以上 -> 上下文区: label 格 -> 同排右侧值格, 进 header_fields,
#                 目录 context 纯字典查找填充 (空值格=覆盖写, 带占位文本=
#                 inline 替换), 与数据区互不干扰;
#   表头行以下 -> 数据区: 按行分组还原 "字段 -> 规定值/实测值/合格率" 结构,
#                 进 rows, 数据简表按 行标签+列名 填充;
#   全表       -> by_label: 所有带 label 的字段 (兜底和通用查找)。
# uuid 由 addr 确定性生成 (uuid5), 同表重建一致; cells 表携带每格几何/
# 占位文本/mode, 回填渲染只查 cells[uuid]。占位覆盖/无值格/缺列等边界
# 构建期全部进 report, 不等到运行期。
#
# 填充语义 (cells[uuid] 的 mode):
#   overwrite  text 为空, 直接写入 value
#   replace    整格替换 (占位文字视为示例值, 如 工程部位 的 "土方路基")
#   append     值追加在占位文本之后 (标签与值同格, 如 抽检时间)
#   inline     替换占位文本片段 (同行冒号内联, 如 编号: JB12-...)
# ---------------------------------------------------------------------------

_INDEX_NS = None  # uuid5 namespace, 延迟生成


def build_index(data):
    """按上述契约生成 data['index']。依赖 extract_fields/assign_roles
    的配对结果 (fields) 与网格几何, 在 build_form_view 之后调用。"""
    import re
    import uuid
    global _INDEX_NS
    if _INDEX_NS is None:
        _INDEX_NS = uuid.UUID('7d4f2a90-3b1c-5e6a-9f08-2d5b8c4a7e91')

    def norm(s):
        return re.sub(r'\s+', '', str(s or ''))

    cells = data['cells']
    by_grid = {(c['row'], c['col']): c for c in cells}
    header_row = min((c['row'] for c in cells if c.get('role') == 'header'),
                     default=None)

    # 同 addr 的游离文字按视觉行拼接 (一个原生格常被拆成多个 seg,
    # 如 A3 = "监理单位：..." + "JB12-4.2.2-"), 拼接后占位提取才完整
    merged_text = {}
    for grp in _group_free_text_lines(data.get('texts') or []):
        by_a = {}
        for it in grp['items']:
            t = data['texts'][it['idx']]
            if t.get('addr'):
                by_a.setdefault(t['addr'], []).append(it)
        for a, items in by_a.items():
            merged_text[a] = ''.join(
                data['texts'][it['idx']]['text']
                for it in sorted(items, key=lambda i: i['x']))

    idx_cells, report = {}, []
    _noaddr = {}  # id(cell) -> uuid (无 addr 的格, 每次运行不同)

    def ref(c):
        """cell -> uuid, 并登记进 cells 表 (几何/文本/占位)。"""
        a = c.get('addr')
        if a:
            u = str(uuid.uuid5(_INDEX_NS, a))
        else:
            u = _noaddr.setdefault(id(c), str(uuid.uuid4()))
        if u not in idx_cells:
            e = {'addr': a, 'x': round(c['x'], 1), 'y': round(c['y'], 1),
                 'w': round(c.get('w') or 0, 1),
                 'h': round(c.get('h') or 0, 1)}
            t = merged_text.get(a, (c.get('text') or '').strip())
            if t:
                e['text'] = t
            # 字号/字体/颜色: PDF 文字层 (lines, 实际打印样式) > 顶层原生
            # 属性 (docx/xlsx 简化排版通路) > openpyxl 原生字体 (空值格
            # 唯一来源) > 无; 对齐只有原生侧可得。
            # 渲染层按此复原字体 (build_template -> style.fontFamily 等)
            ln0 = (c.get('lines') or [{}])[0] or {}
            size = c.get('size') or ln0.get('size') or c.get('font_size')
            bold = (c.get('bold') or ln0.get('bold')
                    or c.get('font_bold'))
            family = (c.get('font') or ln0.get('font')
                      or c.get('font_name'))
            color = (c.get('color') or ln0.get('color')
                     or c.get('font_color'))
            align = c.get('align') or c.get('font_align')
            if size:
                e['size'] = round(float(size), 2)
            if bold:
                e['bold'] = True
            if family:
                e['font'] = str(family)
            if color:
                e['color'] = str(color)
            if align:
                e['align'] = align
            idx_cells[u] = e
        return u

    def mode_of(u, f):
        """按 fields 的配对 mode 标注定填充方式; 占位文本进报告。

        append 型以 prefix 为占位: 值插在标签之后, 原标签后的其余
        内容 (如共享格里的下一字段) 保留。
        """
        if f.get('inline'):
            mode = 'inline'
        elif f.get('replace'):
            mode = 'replace'
        elif f.get('append'):
            mode = 'append'
        else:
            mode = 'overwrite'
        text = idx_cells[u].get('text')
        if mode == 'append' and f.get('prefix'):
            placeholder = f['prefix']
        elif mode == 'inline' and text:
            # 占位 = 最后一个冒号之后的值文本 (整行共享格/整格内联均适用)
            tail = re.split(r'[:：]', text)[-1]
            placeholder = tail.strip() or text
        elif text:
            placeholder = text
        else:
            placeholder = None
        idx_cells[u]['mode'] = mode
        if placeholder:
            idx_cells[u]['placeholder'] = placeholder
            report.append({
                'level': 'info', 'type': 'placeholder_fill',
                'field': f['name'], 'cell': u, 'addr': idx_cells[u]['addr'],
                'placeholder': placeholder, 'mode': mode,
                'message': f'"{f["name"]}" 值格含占位文本 '
                           f'"{placeholder}", 按 {mode} 填充'})
        return mode

    header_fields, by_label, rows = {}, {}, []

    def add_label(c, value_u):
        t = norm((c.get('text') or '').replace('\n', ''))
        if not t:
            return
        ent = by_label.setdefault(t, [])
        u = ref(c)
        if any(e['cell'] == u for e in ent):
            return
        e = {'cell': u}
        if value_u:
            e['value'] = value_u
        ent.append(e)

    if header_row is None:
        report.append({'level': 'warn', 'type': 'no_header',
                       'message': '未识别到表头行, header_fields/rows 为空, '
                                  '仅 by_label 可用'})
    # label 可能是格子也可能是游离文字 (监理单位/编号), 游离文字无
    # 行号, 上下文区用 y 坐标与表头行比较定区
    all_ents = list(cells) + list(data.get('texts') or [])

    def find_ent(addr):
        return next((c for c in all_ents if c.get('addr') == addr), None)

    if header_row is not None:
        hdr_y = min(c['y'] for c in cells if c['row'] == header_row)
        # ---- 上下文区: 表头行以上的 fields (label -> 同排右侧值格) ----
        for f in data.get('fields') or []:
            label_ent = find_ent(f.get('label'))
            if label_ent is None or label_ent.get('y', 0) >= hdr_y:
                continue  # 表尾/表内字段 (检查结论/审核日期等) 不进 header_fields
            if not f.get('value') and not (f.get('inline')
                                           and f.get('prefix')):
                report.append({'level': 'warn', 'type': 'no_value_cell',
                               'field': f['name'], 'addr': f.get('label'),
                               'message': f'"{f["name"]}" 未配到值格, '
                                          f'仅进 by_label'})
                add_label(label_ent, None)
                continue
            value_ent = find_ent(f.get('value'))
            if value_ent is None:
                value_ent = label_ent   # 标签与值同格 (内联前缀型), 追加填充
                f = {**f, 'inline': False, 'append': True}
            u = ref(value_ent)
            key = norm(f['name'])

            def _reg_label():
                # 以字段名登记 by_label (label 全文可能含多字段, 如 A3 行)
                ent = by_label.setdefault(key, [])
                if not any(x['cell'] == ref(label_ent) for x in ent):
                    x = {'cell': ref(label_ent)}
                    if u != ref(label_ent):
                        x['value'] = u
                    ent.append(x)

            if key in header_fields:
                report.append({'level': 'warn', 'type': 'dup_field',
                               'field': key,
                               'message': f'上下文区重名字段 "{key}", 后者 '
                                          f'仅进 by_label'})
                _reg_label()
                continue
            header_fields[key] = {'cell': u}
            header_fields[key]['mode'] = mode_of(u, f)
            _reg_label()

        # ---- 数据区: 表体行 -> 字段 -> 各列 (规定值/实测值/合格率) ----
        headers = {}
        for c in cells:
            if c['row'] == header_row and (c.get('text') or '').strip():
                headers.setdefault(c['col'],
                                   (c.get('text') or '').replace('\n', '')
                                   .strip())
        table_rows = _table_rows_set(cells, header_row)
        counts = {}
        for c in cells:
            if c['row'] in table_rows and (c.get('text') or '').strip():
                counts[c['col']] = counts.get(c['col'], 0) + 1
        if counts and headers:
            label_col = max(counts, key=counts.get)
            for r in sorted(table_rows):
                kc = by_grid.get((r, label_col))
                kt = (kc.get('text') or '').replace('\n', '').strip() \
                    if kc else ''
                if not kt:
                    continue  # 专用填写行/备用空行
                row_doc = {'key': kt, 'key_cell': ref(kc)}
                add_label(kc, None)
                for hc, hname in headers.items():
                    if hc == label_col:
                        continue
                    target = by_grid.get((r, hc))
                    if target is None or not (target.get('text') or '') \
                            .strip():
                        target = None
                        for rr in (r + 1, r):
                            tc = by_grid.get((rr, hc))
                            if tc is None \
                                    or (tc.get('text') or '').strip() \
                                    or not tc.get('addr'):
                                continue
                            if rr == r + 1:
                                lbl = by_grid.get((rr, label_col))
                                if lbl is not None \
                                        and (lbl.get('text') or '').strip():
                                    continue
                            target = tc
                            break
                    if target is None or not target.get('addr'):
                        report.append({'level': 'warn',
                                       'type': 'missing_col',
                                       'row': kt, 'col': hname,
                                       'message': f'行 "{kt}" 缺列 '
                                                  f'"{hname}" 的值格'})
                        continue
                    col = {'cell': ref(target)}
                    if (target.get('text') or '').strip():
                        col['text'] = target['text']
                    row_doc.setdefault('columns', {})[hname] = col
                if row_doc.get('columns'):
                    rows.append(row_doc)
                else:
                    report.append({'level': 'info',
                                   'type': 'no_fillable_col',
                                   'row': kt,
                                   'message': f'行 "{kt}" 没有可填列 '
                                              f'(如检查要求说明行), 不进 rows'})

    # ---- by_label 兜底: 表尾 label 角色格 (检查结论/审核人 等) ----
    registered = {e['cell'] for lst in by_label.values() for e in lst}
    for c in cells:
        if c.get('role') != 'label' or not c.get('addr'):
            continue
        u = ref(c)
        if u in registered:
            continue
        right = by_grid.get((c['row'], c['col'] + c.get('colspan', 1)))
        value_u = None
        if right is not None and right.get('addr'):
            value_u = ref(right)
            if (right.get('text') or '').strip():
                idx_cells[value_u]['mode'] = 'replace'
                idx_cells[value_u]['placeholder'] = \
                    (right.get('text') or '').strip()
                report.append({
                    'level': 'info', 'type': 'placeholder_fill',
                    'field': norm(c.get('text')), 'cell': value_u,
                    'addr': right['addr'],
                    'placeholder': (right.get('text') or '').strip(),
                    'mode': 'replace',
                    'message': f'"{norm(c.get("text"))}" 值格含占位文本, '
                               f'按 replace 填充'})
        add_label(c, value_u)

    # fields 里尚未登记的标签 (审核日期/监理单位/编号 等非 label 角色,
    # label 可能是格子也可能是游离文字)
    for f in data.get('fields') or []:
        label_ent = find_ent(f.get('label'))
        if label_ent is None:
            continue
        u = ref(label_ent)
        if u in {e['cell'] for lst in by_label.values() for e in lst}:
            continue
        value_u = None
        value_ent = find_ent(f.get('value'))
        if value_ent is not None and f.get('value') != f.get('label'):
            value_u = ref(value_ent)
        elif value_ent is not None:
            value_u = u  # label/value 同格 (inline/append)
        # 以字段名 (而非 label 格全文) 登记, 游离文字 label 的全文
        # 可能含多个字段 (如 A3 一整行)
        key = norm(f['name'])
        ent = by_label.setdefault(key, [])
        if not any(x['cell'] == u for x in ent):
            x = {'cell': u}
            if value_u:
                x['value'] = value_u
            ent.append(x)

    data['index'] = {
        'meta': {'title': (data.get('form') or {}).get('meta', {})
                 .get('title'),
                 'dpi': data.get('dpi'), 'width': data.get('width'),
                 'height': data.get('height'),
                 'header_row': header_row,
                 'norm': '去除所有空白字符'},
        'header_fields': header_fields,
        'rows': rows,
        'by_label': by_label,
        'cells': idx_cells,
        'report': report,
    }
    print(f"index: header_fields {len(header_fields)} 个, rows {len(rows)}"
          f" 行, by_label {len(by_label)} 项, 报告 {len(report)} 条")


# ---------------------------------------------------------------------------
# 前端代码生成 (数据驱动: data.json + index.html)
# ---------------------------------------------------------------------------

def generate_frontend(data, out_dir):
    """把标准 JSON 落成前端产物:

    form.json   完整 JSON (含 form 可索引视图), 转换的直接产物,
                建索引/回填/还原都以它为准
    data.json   同一份内容的原名, 兼容已有消费方
    data.js     同一份数据包成 window.TABLE_DATA, 供 index.html 以 file:// 打开
    template.json  前端 Template 结构 (背景图 + 字段覆盖层), Vue 复现层消费,
                由 build_template 生成, 见 build_template docstring
    index.html  纯渲染层: 读数据, 为每个格子生成绝对定位的 div,
                几何与原图像素一一对应, 边框只在原图有线处绘制。
    """
    json_text = json.dumps(data, ensure_ascii=False, indent=2)
    for name in ('form.json', 'data.json'):
        with open(os.path.join(out_dir, name), 'w', encoding='utf-8') as f:
            f.write(json_text)
    with open(os.path.join(out_dir, 'data.js'), 'w', encoding='utf-8') as f:
        f.write(f'window.TABLE_DATA = {json_text};\n')

    index_path = os.path.join(out_dir, 'index.html')
    with open(index_path, 'w', encoding='utf-8') as f:
        f.write(INDEX_HTML)
    return index_path


INDEX_HTML = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>还原表格</title>
<style>
  body { background: #f5f5f5; font-family: "Microsoft YaHei", sans-serif; }
  .sheet { position: relative; background: #fff; margin: 24px auto; }
  .cell { position: absolute; box-sizing: border-box; padding: 4px 8px;
          white-space: pre-line; overflow: hidden;
          display: flex; align-items: center; }
  .freetext { position: absolute; white-space: nowrap; line-height: 1; }
</style>
</head>
<body>
  <div class="sheet" id="sheet"></div>
  <script src="data.js"></script>
  <script>
    const data = window.TABLE_DATA;
    const sheet = document.getElementById('sheet');
    sheet.style.width = data.width + 'px';
    sheet.style.height = data.height + 'px';
    const pxPerPt = (data.dpi || 300) / 72;
    const rules = data.rules || [];
    for (const c of data.cells) {
      const div = document.createElement('div');
      div.className = 'cell';
      div.style.left = c.x + 'px';
      div.style.top = c.y + 'px';
      div.style.width = c.w + 'px';
      div.style.height = c.h + 'px';
      if (!rules.length)  // 有矢量线段时由 rules 负责描线
        for (const side of ['top', 'bottom', 'left', 'right'])
          div.style['border' + side[0].toUpperCase() + side.slice(1)] =
            c.borders[side] ? '2px solid #333' : '0';
      div.textContent = c.text;
      sheet.appendChild(div);
    }
    // 电子档原始矢量线段: 精确坐标 + 原线宽, 每根线只画一次
    for (const rl of rules) {
      const div = document.createElement('div');
      const w = Math.abs(rl.x1 - rl.x0), h = Math.abs(rl.y1 - rl.y0);
      div.style.position = 'absolute';
      div.style.left = Math.min(rl.x0, rl.x1) + 'px';
      div.style.top = Math.min(rl.y0, rl.y1) + 'px';
      if (h > w) {  // 竖线
        div.style.width = Math.max(1, rl.width) + 'px';
        div.style.height = h + 'px';
      } else {      // 横线
        div.style.width = w + 'px';
        div.style.height = Math.max(1, rl.width) + 'px';
      }
      div.style.background = '#000';
      sheet.appendChild(div);
    }
    // 表格线框外的游离文字 (页眉/标题/页码等): x,y 为基线坐标(px), size 单位 pt
    for (const t of (data.texts || [])) {
      const div = document.createElement('div');
      div.className = 'freetext';
      const fs = t.size * pxPerPt;
      div.style.left = t.x + 'px';
      div.style.top = (t.y - fs * 0.85) + 'px';
      div.style.fontSize = fs + 'px';
      div.style.color = t.color || '#000';
      if (t.bold || t.role === 'title' || t.role === 'heading')
        div.style.fontWeight = 'bold';
      div.dataset.role = t.role || 'body';
      div.textContent = t.text;
      sheet.appendChild(div);
    }
  </script>
</body>
</html>
"""


# ---------------------------------------------------------------------------
# 可视化
# ---------------------------------------------------------------------------

def draw_overlay(img, h_pos, v_pos, cells, out_path, blocks=None):
    vis = img.copy()
    for b in blocks or []:
        color = (200, 0, 200) if b['type'] == 'text' else (0, 165, 255)
        cv2.rectangle(vis, (b['x'], b['y']),
                      (b['x'] + b['w'], b['y'] + b['h']), color, 2)
    for y in h_pos:
        cv2.line(vis, (0, y), (vis.shape[1] - 1, y), (0, 255, 0), 1)
    for x in v_pos:
        cv2.line(vis, (x, 0), (x, vis.shape[0] - 1), (255, 0, 0), 1)
    for c in cells:
        cv2.rectangle(vis, (c['x'], c['y']),
                      (c['x'] + c['w'], c['y'] + c['h']), (0, 0, 255), 1)
        cv2.putText(vis, f"{c['row']},{c['col']}",
                    (c['x'] + 3, c['y'] + 15),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.4, (0, 0, 255), 1)
    cv2.imwrite(out_path, vis)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def stitch_page_data(pages):
    """宽表被分页时, 把后续页横向拼接到第一页右边。

    横向延续的判定: 该页格子的 y 行位与第一页表格行位高度重合
    (同一批行被切到右边); 对不上说明是纵向分页(长表), 不拼。
    """
    base = pages[0]
    if len(pages) == 1 or not base['cells']:
        return base
    ys0 = sorted({c['y'] for c in base['cells']})
    margin = min([c['x'] for c in base['cells']]
                 + [r['x0'] for r in base['rules']])
    out = dict(base)
    out['cells'] = list(base['cells'])
    out['texts'] = list(base['texts'])
    out['rules'] = list(base['rules'])
    right = max([c['x'] + c['w'] for c in base['cells']]
                + [r['x1'] for r in base['rules']])
    stitched = 0
    row_off = (max(c['row'] for c in base['cells']) + 1
               if base['cells'] else 0)  # 行号全局唯一化, 供网格邻接判定
    for pg in pages[1:]:
        if not pg['cells']:
            break
        ys = [c['y'] for c in pg['cells']]
        hit = sum(1 for y in ys if any(abs(y - y0) <= 6 for y0 in ys0))
        if hit < len(ys) * 0.7:
            break  # 行位对不上 -> 纵向分页, 不拼
        # 分页按列边界切: 接缝 = 已拼内容的最右缘, 该页内容左 margin 对齐过来
        left = min([c['x'] for c in pg['cells']]
                   + [r['x0'] for r in pg['rules']])
        dx = right - left
        for c in pg['cells']:
            c = dict(c)
            c['x'] += dx
            c['row'] += row_off
            if 'lines' in c:
                c['lines'] = [dict(l, x=l['x'] + dx) if 'x' in l else l
                              for l in c['lines']]
            out['cells'].append(c)
        for t in pg['texts']:
            out['texts'].append(dict(t, x=t['x'] + dx))
        for r in pg['rules']:
            out['rules'].append({**r, 'x0': r['x0'] + dx,
                                 'x1': r['x1'] + dx})
        right = max(right, max([c['x'] + c['w'] for c in out['cells']
                                if c['x'] >= dx]
                               or [right]))
        row_off += max(c['row'] for c in pg['cells']) + 1
        stitched += 1
    if stitched:
        out['width'] = right + margin  # 右侧补一个对称 margin
        out['height'] = max(p['height'] for p in pages[:stitched + 1])
        print(f'宽表分页拼接: {stitched + 1} 页横向合并')
    return out


def scale_data(data, s):
    """对标准 JSON 做纯数学等比缩放 (坐标/字号都乘 s), 相对位置不变。"""
    if abs(s - 1.0) < 1e-6:
        return data
    out = dict(data)
    cells = []
    for c in data['cells']:
        c = dict(c)
        for k in ('x', 'y', 'w', 'h'):
            c[k] = c[k] * s
        if 'size' in c:
            c['size'] *= s
        if 'lines' in c:
            ls = []
            for l in c['lines']:
                l = dict(l)
                for k in ('x', 'y', 'w', 'h'):
                    if k in l:
                        l[k] *= s
                if 'size' in l:
                    l['size'] *= s
                ls.append(l)
            c['lines'] = ls
        cells.append(c)
    out['cells'] = cells
    out['texts'] = [dict(t, x=t['x'] * s, y=t['y'] * s, size=t['size'] * s,
                         w=t['w'] * s if t.get('w') else t.get('w'))
                    for t in data.get('texts', [])]
    out['rules'] = [dict(r, x0=r['x0'] * s, y0=r['y0'] * s,
                         x1=r['x1'] * s, y1=r['y1'] * s,
                         width=r['width'] * s)
                    for r in data.get('rules', [])]
    out['width'] = data['width'] * s
    out['height'] = data['height'] * s
    return out


def extract_page(input_path, page, args, diagnostics=True):
    """单页处理: 线检测 -> 格子 -> 文字 -> 标准 JSON data。

    diagnostics=True 时输出 overlay.png / lines.json 和统计信息。
    """
    img = load_image(input_path, page=page, dpi=args.dpi)
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)

    bw = binarize(gray)
    horiz, vert = extract_line_masks(bw, h_ratio=args.h_ratio,
                                     v_ratio=args.v_ratio)
    # 交点判定和格子边判定用接合断口后的掩码, 避免断线处的误判
    horiz_c = close_gaps(horiz, 'h')
    vert_c = close_gaps(vert, 'v')
    h_pos = find_line_positions(horiz_c, 'h', min_span=args.h_min_span)
    v_pos = find_line_positions(vert_c, 'v', min_span=args.v_min_span)
    # 表格线的两端必然落在对向线上; 两端悬空的(页眉下划线等)滤掉
    h_pos = filter_dangling_lines(h_pos, horiz_c, 'h', v_pos)
    v_pos = filter_dangling_lines(v_pos, vert_c, 'v', h_pos)

    # 电子档 PDF: OpenCV 检测(像素)与原始矢量线段(get_drawings)合并——
    # 掩码按位或、线坐标合并聚类, 再统一过悬空过滤。纯 OpenCV 在表格只占
    # 页面一角时容易丢外框线(邻线缺失引发连锁过滤), 矢量线段正好补上
    is_pdf = os.path.splitext(input_path)[1].lower() == '.pdf'
    rules = []
    if is_pdf and not args.force_ocr:
        rules = extract_pdf_rules(input_path, page, args.dpi)
        if rules:
            horiz_r, vert_r, h_r, v_r = rules_to_masks(rules, gray.shape)
            horiz_c = cv2.bitwise_or(horiz_c, horiz_r)
            vert_c = cv2.bitwise_or(vert_c, vert_r)
            h_pos = _merge_positions(h_pos, h_r)
            v_pos = _merge_positions(v_pos, v_r)
            h_pos = filter_dangling_lines(h_pos, horiz_c, 'h', v_pos)
            v_pos = filter_dangling_lines(v_pos, vert_c, 'v', h_pos)

    intersections = compute_intersections(h_pos, v_pos, horiz_c, vert_c)
    cells = find_cells(h_pos, v_pos, horiz_c, vert_c)

    # 版面分析: 页面内容边界 + 文字块/图片块; 每格做横向投影行切分,
    # 行切分结果用于约束 OCR 文字还原时的字号和位置
    content, blocks, bounds = analyze_layout(bw, horiz_c, vert_c)
    for c in cells:
        c['text_lines'] = split_text_lines(
            content, (c['x'] + 4, c['y'] + 4, c['w'] - 8, c['h'] - 8))

    texts = {}
    free_texts = []
    if not args.no_ocr:
        need_ocr = True
        if is_pdf and not args.force_ocr:
            # 探测文字层: 有 -> 电子档, 直接提取; 无 -> 扫描件, 回退 OCR
            pdf_texts, free_texts = extract_pdf_text(input_path, page,
                                                     args.dpi, cells)
            if pdf_texts or free_texts:
                texts = pdf_texts or {}
                need_ocr = False
                if diagnostics:
                    roles = {}
                    for t in free_texts:
                        roles[t['role']] = roles.get(t['role'], 0) + 1
                    print(f'文字层探测: 电子档, 直接提取 (格内 {len(texts)}/'
                          f'{len(cells)} 格 + 格外 {len(free_texts)} 行 '
                          f'{roles} + 矢量线段 {len(rules)} 条)')
            else:
                print('文字层探测: 扫描件(无文字层), 回退到 OCR')
        if need_ocr:
            name, backend = _pick_ocr_backend(args)
            print(f'OCR 后端: {name} (逐格请求)')
            ocr_res = backend(img, cells)
            if ocr_res is None:
                print('未安装 rapidocr-onnxruntime, 跳过文字识别'
                      ' (pip install rapidocr-onnxruntime);'
                      ' 或用 --baidu-ak/--baidu-sk'
                      ' / --paddle-url/--paddle-token 走云端 OCR')
            else:
                texts = ocr_res
                # 投影行切分对位: 约束每行文字还原时的字号和位置
                _attach_line_pos(cells, texts)
                # 表格线框外的文字块也 OCR, 按行高估算字号作为游离文字
                free_texts = ocr_free_blocks(img, blocks, cells, content,
                                             backend, args.dpi)
                if free_texts:
                    tmp = [{'size': t['size'], 'cx': t.pop('cx')}
                           for t in free_texts]
                    _classify_line_roles(tmp, img.shape[1])
                    for t, r in zip(free_texts, tmp):
                        t['role'] = r['role']
                    print(f'格外游离文字 {len(free_texts)} 行 (OCR)')
        if texts and diagnostics:
            print(f'文字提取完成, {len(texts)}/{len(cells)} 个格子有文字')
            for c in cells:
                t = texts.get((c['row'], c['col']))
                if t:
                    print(f'  格[{c["row"]},{c["col"]}]: {t}')

    data = build_table_data(cells, texts, horiz_c, vert_c, args.dpi,
                            (img.shape[1], img.shape[0]), free_texts, rules)

    if diagnostics:
        n_txt = sum(1 for b in blocks if b['type'] == 'text')
        print(f'版面分析: 内容边界 {bounds}, 文字块 {n_txt} 个, '
              f'图片块 {len(blocks) - n_txt} 个')
        overlay_path = os.path.join(args.out, 'overlay.png')
        draw_overlay(img, h_pos, v_pos, cells, overlay_path, blocks)
        result = {
            'image_size': {'width': img.shape[1], 'height': img.shape[0]},
            'horizontal_lines': h_pos,
            'vertical_lines': v_pos,
            'intersections': intersections,
            'cells': cells,
            'layout': {'content_bounds': bounds, 'blocks': blocks},
        }
        for c in cells:
            c['text'] = texts.get((c['row'], c['col']), '')
        json_path = os.path.join(args.out, 'lines.json')
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f'横线 {len(h_pos)} 条, 竖线 {len(v_pos)} 条, '
              f'交点 {len(intersections)} 个, 格子 {len(cells)} 个')
        print(f'结果已写入 {json_path}')
        print(f'可视化: {overlay_path}')
    return data


def convert_document(input_path, out_dir, args):
    """docx/xlsx 完整转换流水线 (main 与 server 共用)。

    优先用本机 WPS/Office 排版引擎 100% 原样转 PDF, 走电子档通路;
    失败才回退到 python-docx/openpyxl 的简化排版。
    依次产出: form.json/data.json/data.js/index.html (generate_frontend)
    和 rebuilt.pdf。返回 (data, office_pdf, pdf_path): office_pdf 是
    排版引擎转出的原样 PDF (无 WPS/Office 时为 None)。
    """
    ext = os.path.splitext(input_path)[1].lower()
    conv = None if args.no_office else office_to_pdf(input_path, out_dir)
    if conv:
        fitz = _import_pymupdf()
        doc = fitz.open(conv)
        n_pages = len(doc)
        doc.close()
        pages = [extract_page(conv, p, args, diagnostics=(p == 0))
                 for p in range(n_pages)]
        if ext == '.xlsx':
            # 语义对齐: 逐页把探测网格锚定到原生 workbook,
            # 给格子/游离文字赋稳定地址 addr (stitch 只平移 x, 先行列号不变)
            try:
                import openpyxl
                wb_f = openpyxl.load_workbook(input_path)
                wb_v = openpyxl.load_workbook(input_path, data_only=True)
                for pg in pages:
                    augment_xlsx_semantics(pg, wb_f, wb_v)
            except ImportError:
                print('未安装 openpyxl, 跳过语义对齐')
            # 宽表被分页: 先横向拼回一整张, 再在数据层等比缩放到
            # A4 页宽 (纯数学缩放, 相对位置严格不变)
            data = stitch_page_data(pages)
            if data['width'] > pages[0]['width']:
                s = pages[0]['width'] / data['width']
                data = scale_data(data, s)
                print(f'按 A4 页宽等比缩放到 {s:.0%}')
        else:
            data = pages[args.page] if args.page < n_pages else pages[0]
    else:
        data = (docx_to_data(input_path) if ext == '.docx'
                else xlsx_to_data(input_path))
    assign_roles(data)
    extract_fields(data)
    build_form_view(data)
    build_index(data)
    generate_frontend(data, out_dir)
    pdf_path = json_to_pdf(data, os.path.join(out_dir, 'rebuilt.pdf'))
    return data, conv, pdf_path


def _font_family(name):
    """归一化字体名: 去 PDF 子集前缀 (ABCDEE+SimSun), 中文名映射英文，
    对齐前端 fonts.ts 的 documentFonts 候选。"""
    if not name:
        return None
    name = str(name).split('+')[-1].strip()
    return _FONT_ALIASES.get(name, name)


_FONT_ALIASES = {
    '宋体': 'SimSun', '黑体': 'SimHei', '微软雅黑': 'Microsoft YaHei',
    '仿宋': 'FangSong', '楷体': 'KaiTi', '等线': 'DengXian',
    'ArialMT': 'Arial', 'TimesNewRomanPSMT': 'Times New Roman',
}


def _field_style(e, size):
    """index.cells 实体 -> 前端 FieldStyle (字体/字号/颜色/对齐)。"""
    style = {'fontSizePt': size,
             'horizontalAlign': e.get('align') or 'center',
             'verticalAlign': 'center'}
    fam = _font_family(e.get('font'))
    if fam:
        style['fontFamily'] = fam
    if e.get('color'):
        style['color'] = e['color']
    if e.get('bold'):
        style['bold'] = True
    return style


def build_template(data, out_dir, source_paths=(), fallback_pdf=None,
                   name=None, bg_url='background.png'):
    """生成前端 Template JSON (template.json) + 背景图 background.png。

    前端复现采用"背景图 + 字段覆盖层"混合方案 (frontend TemplatePage.vue):
    - 背景图: pymupdf 把排版引擎的原样 PDF (office 转换件; 几何不匹配时
      退到 rebuilt.pdf) 按 2px/pt 栅格化成 background.png,
      "原样张"视觉上由它保证;
    - 字段层: OpenCV 探测网格的每一格变成一个 Field (与背景表格线严格
      对齐), rect 为占页宽/页高的比例坐标, uuid = uuid5(addr) 与
      index.cells 同一 uuid —— header_fields/rows/by_label 的 value
      引用同一 uuid, 回填定位零换算; index 解析好的样式/填充语义按
      uuid 挂载, 识别文字进 label (不做覆盖层重画, 视觉以背景图为
      唯一真相, 覆盖层只在填值时出现)。
    page.width/height 用背景图像素尺寸 (2px/pt 栅格的宽/高), 不是 PDF 点:
    前端字号公式 fontSizePt*2/page.width*100cqw (TemplatePage.vue) 和
    PrintPreview 的 page.width/2*25.4/72 (mm) 都按这个约定消费。
    无 index 的通路 (图片/PDF) 退化为逐格生成 Field。
    """
    import uuid as _uuid
    fitz = _import_pymupdf()
    dpi = data.get('dpi') or 300
    pt_w = data['width'] * 72.0 / dpi
    pt_h = data['height'] * 72.0 / dpi

    # 背景源: 逐个候选试, 几何 (页面点大小) 与 data 一致才采用
    bg_path = os.path.join(out_dir, 'background.png')
    used_src = None
    bg_wh = None
    for cand in list(source_paths) + ([fallback_pdf] if fallback_pdf else []):
        if not cand or not os.path.exists(cand):
            continue
        doc = None
        try:
            doc = fitz.open(cand)
            if len(doc) >= 1:
                p = doc[0]
                if abs(p.rect.width - pt_w) <= max(2.0, pt_w * 0.02) and \
                        abs(p.rect.height - pt_h) <= max(2.0, pt_h * 0.02):
                    pix = p.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    pix.save(bg_path)
                    used_src = cand
                    bg_wh = (pix.width, pix.height)
        except Exception:
            pass
        finally:
            if doc is not None:
                doc.close()
        if used_src:
            break

    # 字段层: OpenCV 探测网格的每一格 -> Field, 与背景表格线严格对齐;
    # uuid = uuid5(addr), 与 index.cells 同一 uuid, 回填定位零换算;
    # index 解析好的样式 (字号/字体/颜色/对齐/合并文本/填充语义) 按 uuid
    # 挂到对应格上, 没进 index 的普通格直接用自身文字层的样式。
    # 识别文字一律进 label (不进 textOverride, 覆盖层只在运行时填值出现)
    ns = _INDEX_NS or _uuid.UUID('7d4f2a90-3b1c-5e6a-9f08-2d5b8c4a7e91')
    idx_cells = ((data.get('index') or {}).get('cells')) or {}
    W, H = data['width'], data['height']
    fields = []
    seen = set()

    def _emit(u, x, y, w, h, ent, text, addr):
        w = min(w, W)
        h = min(h, H)
        x = max(0.0, min(x, W - w))
        y = max(0.0, min(y, H - h))
        f = {'fieldUuid': u,
             'rect': {'x': round(x / W, 6), 'y': round(y / H, 6),
                      'w': round(w / W, 6), 'h': round(h / H, 6)},
             'style': _field_style(ent, round(ent.get('size') or 10, 2))}
        if text:
            f['label'] = text
        if addr:
            f['source'] = {'range': addr}
        if ent.get('mode'):
            f['mode'] = ent['mode']
            if ent.get('placeholder'):
                f['placeholder'] = ent['placeholder']
        fields.append(f)
        seen.add(u)

    for c in data.get('cells') or []:
        a = c.get('addr')
        u = str(_uuid.uuid5(ns, a)) if a else str(
            _uuid.uuid5(ns, f"#{c.get('row')},{c.get('col')}"))
        e = idx_cells.get(u) or {}
        # index 合并文本 (同 addr 游离文字按视觉行拼接) 优先于格内 text;
        # 样式: index 解析结果覆盖格内原始属性
        text = e.get('text') or (c.get('text') or '').strip() or None
        ent = dict(c)
        ent.update(e)
        _emit(u, c['x'], c['y'], c['w'], c['h'], ent, text, a)
    # index 里的游离文字实体 (监理单位/编号等, 无格子) 也生成字段,
    # 几何按字号估算
    for u, e in idx_cells.items():
        if u in seen:
            continue
        x, y, w, h = e['x'], e['y'], e['w'], e['h']
        size = e.get('size') or 10
        if (not w or not h) and e.get('text'):
            n = len(''.join(str(e['text']).split()))
            w = w or n * size * dpi / 72.0
            h = h or size * dpi / 72.0 * 1.4
            y = y - h * 0.85
        _emit(u, x, y, w, h, e, e.get('text'), e.get('addr'))

    form_meta = ((data.get('form') or {}).get('meta')) or {}
    tpl_name = (name or form_meta.get('title') or '未命名表单')
    ext = os.path.splitext(os.path.basename(used_src or ''))[1]
    template = {
        'templateUuid': str(_uuid.uuid5(ns, 'template:' + tpl_name)),
        'name': tpl_name,
        'version': 1,
        'sourceFormat': (ext.lstrip('.') if used_src
                         else os.path.splitext(os.path.basename(
                             fallback_pdf or ''))[1].lstrip('.') or 'data'),
        'pages': [{
            'pageUuid': str(_uuid.uuid5(ns, 'page:' + tpl_name)),
            'pageNo': 1,
            # 背景像素尺寸 (无背景时按 2px/pt 折算, 保持同一约定)
            'width': bg_wh[0] if bg_wh else round(pt_w * 2, 2),
            'height': bg_wh[1] if bg_wh else round(pt_h * 2, 2),
            'background': bg_url if used_src else '',
            'fields': fields,
        }],
    }
    tpl_path = os.path.join(out_dir, 'template.json')
    with open(tpl_path, 'w', encoding='utf-8') as f:
        json.dump(template, f, ensure_ascii=False, indent=2)
    print(f'前端模板: {tpl_path} (字段 {len(fields)} 个, 背景 '
          f'{os.path.basename(used_src) if used_src else "无"})')
    return tpl_path


def main():
    ap = argparse.ArgumentParser(
        description='图片/PDF/Word/Excel 表格提取与还原')
    ap.add_argument('input', help='图片、PDF、Word(.docx) 或 Excel(.xlsx) 路径')
    ap.add_argument('-o', '--out', default='grid_out', help='输出目录')
    ap.add_argument('--page', type=int, default=0, help='PDF 页码 (从 0 开始)')
    ap.add_argument('--dpi', type=int, default=300, help='PDF 渲染 dpi')
    ap.add_argument('--h-ratio', type=float, default=0.1,
                    help='横线形态学核占图宽比例 (默认 0.1)')
    ap.add_argument('--v-ratio', type=float, default=0.03,
                    help='竖线形态学核占图高比例 (默认 0.03, 竖线常有短段)')
    ap.add_argument('--h-min-span', type=float, default=0.28,
                    help='横线最小覆盖度 (默认 0.28, 用于滤掉页眉下划线等)')
    ap.add_argument('--v-min-span', type=float, default=0.04,
                    help='竖线最小覆盖度 (默认 0.04, 短竖线覆盖度低)')
    ap.add_argument('--no-ocr', action='store_true',
                    help='跳过 OCR, 生成的 HTML 只还原结构不填文字')
    ap.add_argument('--force-ocr', action='store_true',
                    help='PDF 即使有文字层也强制走 OCR')
    ap.add_argument('--no-office', action='store_true',
                    help='docx/xlsx 不调用本机 WPS/Office 转 PDF, '
                         '直接用 python-docx/openpyxl 简化排版')
    ap.add_argument('--paddle-url', default=os.environ.get('PADDLE_OCR_URL'),
                    help='PaddleOCR API 地址 (或设环境变量 PADDLE_OCR_URL), '
                         '在 aistudio.baidu.com/paddleocr/task 获取')
    ap.add_argument('--paddle-token', default=os.environ.get('PADDLE_OCR_TOKEN'),
                    help='PaddleOCR API token (或设环境变量 PADDLE_OCR_TOKEN)')
    ap.add_argument('--baidu-ak', default=os.environ.get('BAIDU_OCR_AK'),
                    help='百度智能云文字识别 API Key (或环境变量 BAIDU_OCR_AK)')
    ap.add_argument('--baidu-sk', default=os.environ.get('BAIDU_OCR_SK'),
                    help='百度智能云文字识别 Secret Key (或环境变量 BAIDU_OCR_SK)')
    ap.add_argument('--baidu-api', default='accurate_basic',
                    choices=['accurate_basic', 'general_basic', 'handwriting'],
                    help='百度 OCR 接口 (默认 accurate_basic 高精度版)')
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)

    ext = os.path.splitext(args.input)[1].lower()
    stem = os.path.splitext(os.path.basename(args.input))[0]
    if ext in ('.docx', '.xlsx'):
        data, office_pdf, pdf_path = convert_document(args.input, args.out,
                                                      args)
        build_template(data, args.out, source_paths=[office_pdf],
                       fallback_pdf=pdf_path, name=stem)
        print(f'格子 {len(data["cells"])} 个, '
              f'段落 {len(data["texts"])} 行, 线段 {len(data["rules"])} 条')
        print(f'还原 PDF: {pdf_path} (数据源: data.json)')
        print(f'前端代码: {os.path.join(args.out, "index.html")} '
              f'(数据: data.json / data.js / template.json)')
        return

    data = extract_page(args.input, args.page, args)
    index_path = pdf_path = None
    if data:
        index_path = generate_frontend(data, args.out)
        pdf_path = json_to_pdf(data, os.path.join(args.out, 'rebuilt.pdf'))
        build_template(data, args.out,
                       source_paths=[args.input] if ext == '.pdf' else [],
                       fallback_pdf=pdf_path, name=stem)
    if pdf_path:
        print(f'还原 PDF: {pdf_path} (数据源: data.json)')
    if index_path:
        print(f'前端代码: {index_path} (数据: data.json / data.js)')


if __name__ == '__main__':
    main()
